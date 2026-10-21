from __future__ import annotations

import asyncio
from datetime import UTC, datetime
from io import BytesIO
import json
from pathlib import Path
import re
from time import perf_counter
from urllib.parse import quote
from uuid import uuid4
from collections import Counter

from dotenv import load_dotenv
from fastapi import BackgroundTasks, Depends, FastAPI, File, Form, Header, HTTPException, Query, Request, UploadFile, WebSocket, WebSocketDisconnect
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse
from starlette.middleware.gzip import GZipMiddleware
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field

load_dotenv(dotenv_path=Path(__file__).resolve().parents[2] / ".env", override=False)

from contracts import (
    AssessmentAnswer,
    CollectionScheduleRequest,
    CollectionStrategyRequest,
    CleanedDataItem,
    DataCollectionRequest,
    EmergingJobDiscoveryRequest,
    GameResultInput,
    ImportRequest,
    JobGraphBuildRequest,
    MatchReport,
    CandidateProfile,
    JobGraph,
    SkillEvidence,
    SkillRequirement,
    ValidationRunRequest,
)
from core.orchestrator import orchestrator
from core.store import store
from core.auth import AuthError, AuthService
from agents.assessment.resume_files import ResumeFileError, extract_resume_text
from agents.assessment.service import AssessmentQuestionError
from agents.assessment.resume_parser import ResumeParseError
from agents.speech import SpeechModelUnavailable, SpeechRecognitionService
from agents.data.collection import CollectionScheduler
from agents.data.browser_sessions import browser_sessions
from agents.data.cookie_sessions import cookie_sessions
from agents.match.report_export import export_match_report
from agents.data.normalization import extract_skills, is_valid_skill
from agents.data.source_policy import RECRUITMENT_SOURCE_IDS, RECRUITMENT_SOURCE_SET, SOURCE_STATE_PATHS
from agents.learning.service import CourseRagService
from agents.learning.models import JobProfile, JobSkill, RouteRequest
from agents.learning.job_catalog import JOB_FAMILIES

app = FastAPI(title="讯飞智聘 API", version="0.1.0")
collection_scheduler = CollectionScheduler(orchestrator.data_collection_agent)
resume_parse_jobs: dict[str, dict] = {}
auth_service = AuthService(store.db_path or "data/talentgraph.db")
speech_service = SpeechRecognitionService()
learning_service = CourseRagService(store.db_path or "data/talentgraph.db")


class CookieSessionRequest(BaseModel):
    cookies: list[dict[str, object]] = Field(min_length=1, max_length=200)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)
app.add_middleware(GZipMiddleware, minimum_size=1000, compresslevel=5)


def _bearer_token(authorization: str | None) -> str | None:
    if not authorization:
        return None
    scheme, _, token = authorization.partition(" ")
    return token.strip() if scheme.casefold() == "bearer" and token.strip() else None


def current_user(authorization: str | None = Header(default=None)) -> dict:
    user = auth_service.resolve_token(_bearer_token(authorization))
    if not user:
        raise HTTPException(status_code=401, detail="请先登录后继续")
    return user


def require_role(*roles: str):
    def dependency(user: dict = Depends(current_user)) -> dict:
        if user["role"] not in roles:
            raise HTTPException(status_code=403, detail="当前账号没有该功能权限")
        return user

    return dependency


def _required_role(path: str, method: str) -> str | None:
    if method == "OPTIONS":
        return None
    enterprise_prefixes = (
        "/api/data/",
        "/api/collection/",
        "/api/validation/",
    )
    if path.startswith(enterprise_prefixes) or path == "/api/evolution/run":
        return "enterprise"
    if path.startswith("/api/emerging-jobs"):
        return "authenticated" if method == "GET" else "enterprise"
    if path == "/api/jobs" and method != "GET":
        return "enterprise"
    if path.startswith("/api/jobs/") and method in {"POST", "PUT", "PATCH", "DELETE"}:
        return "enterprise"
    if path.startswith(("/api/assessment/", "/api/match/")) and method == "POST":
        return "authenticated"
    return None


@app.middleware("http")
async def enforce_role_permissions(request: Request, call_next):
    required = _required_role(request.url.path, request.method)
    if not required:
        return await call_next(request)
    user = auth_service.resolve_token(_bearer_token(request.headers.get("Authorization")))
    if not user:
        return JSONResponse(status_code=401, content={"detail": "请先登录后继续"})
    if required != "authenticated" and user["role"] != required:
        return JSONResponse(status_code=403, content={"detail": "当前账号没有该功能权限"})
    request.state.user = user
    return await call_next(request)


class AssessmentStartRequest(BaseModel):
    candidate_id: str = "candidate_demo"
    target_job_id: str
    resume_text: str | None = None
    resume_parse: dict | None = None


class ResumeParseRequest(BaseModel):
    candidate_id: str | None = None
    resume_text: str
    source_name: str = "resume_text"


class AssessmentFinalizeRequest(BaseModel):
    session_id: str


class MatchReportRequest(BaseModel):
    candidate_id: str
    target_job_id: str


class GapAnalysisRequest(BaseModel):
    candidate_id: str
    target_job_id: str | None = None
    analyze_all: bool = False


class LearningRouteRequest(BaseModel):
    analysis_id: str
    job_id: str | None = None
    current_level: str = Field(default="入门", pattern="^(零基础|入门|进阶)$")
    courses_per_skill: int = Field(default=3, ge=1, le=6)


class ReportExportRequest(BaseModel):
    report_id: str = Field(min_length=1, max_length=128)


class AssessmentRecordRequest(BaseModel):
    record: dict


class AuthLoginRequest(BaseModel):
    email: str
    password: str
    role: str


class AuthRegisterRequest(BaseModel):
    email: str
    password: str
    role: str
    display_name: str
    organization_name: str | None = None


class ApplicationCreateRequest(BaseModel):
    job_id: str
    job_title: str
    company: str
    source_url: str | None = None
    match_score: int | None = Field(default=None, ge=0, le=100)
    organization_id: str | None = None
    snapshot: dict = Field(default_factory=dict)


class ApplicationStatusRequest(BaseModel):
    status: str


class EvolutionRunRequest(BaseModel):
    job_id: str


class DataReviewRequest(BaseModel):
    status: str
    note: str | None = None
    decision: str | None = None
    corrections: dict = Field(default_factory=dict)


class EmergingJobReviewRequest(BaseModel):
    status: str
    note: str | None = None
    corrections: dict = Field(default_factory=dict)


ALLOWED_DOMAINS = {"人工智能", "大数据", "智能系统", "物联网"}


def ensure_recruitment_source(source_platform: str | None) -> None:
    if source_platform is not None and source_platform not in RECRUITMENT_SOURCE_SET:
        raise HTTPException(
            status_code=422,
            detail=f"source_platform must be one of: {', '.join(RECRUITMENT_SOURCE_IDS)}",
        )


def looks_like_salary(value: str | None) -> bool:
    if not value:
        return False
    return bool(re.search(r"\d", value) and re.search(r"万|千|元|k|K|薪|-|~|－", value))


def valid_city(value: str | None) -> bool:
    return bool(value and not looks_like_salary(value) and len(value) <= 20)


def configured_job_cities() -> list[str]:
    strategies = orchestrator.data_collection_agent.source_strategies
    ordered = []
    for source_id in RECRUITMENT_SOURCE_IDS:
        for city in strategies.get(source_id, {}).get("cities", []):
            if city != "全国" and city not in ordered:
                ordered.append(city)
    return ordered


class JobCreateRequest(BaseModel):
    source_id: str | None = None
    title: str
    domain: str
    company: str
    content: str = ""
    city: str | None = None
    district: str | None = None
    salary_min_k: int | None = None
    salary_max_k: int | None = None
    experience_min_years: int | None = None
    experience_max_years: int | None = None
    education: str | None = None
    skill_tags: list[str] = []
    responsibilities: list[str] = []
    source_platform: str = "企业岗位库"
    source_url: str | None = None
    job_status: str = "active"


class JobUpdateRequest(BaseModel):
    title: str | None = None
    domain: str | None = None
    company: str | None = None
    content: str | None = None
    city: str | None = None
    district: str | None = None
    salary_min_k: int | None = None
    salary_max_k: int | None = None
    experience_min_years: int | None = None
    experience_max_years: int | None = None
    education: str | None = None
    skill_tags: list[str] | None = None
    responsibilities: list[str] | None = None
    source_platform: str | None = None
    source_url: str | None = None
    job_status: str | None = None


class EnterpriseJobRequest(BaseModel):
    title: str = Field(min_length=2, max_length=120)
    domain: str
    city: str = Field(min_length=1, max_length=40)
    headcount: int = Field(default=1, ge=1, le=999)
    responsibilities: list[str] = Field(default_factory=list)
    required_skills: list[str] = Field(default_factory=list)
    preferred_skills: list[str] = Field(default_factory=list)
    experience_min_years: int | None = Field(default=None, ge=0, le=40)
    education: str | None = None
    salary_min_k: int | None = Field(default=None, ge=0, le=500)
    salary_max_k: int | None = Field(default=None, ge=0, le=500)
    status: str = "active"


class RecruitmentFeedbackRequest(BaseModel):
    decision: str = Field(pattern="^(pending_review|advance_interview|rejected|offer)$")
    note: str = Field(default="", max_length=4000)
    ratings: dict[str, int] = Field(default_factory=dict)


def _enterprise_job_payload(item: EnterpriseJobRequest, organization_id: str, *, job_id: str | None = None) -> dict:
    now = datetime.now(UTC).isoformat()
    return {
        "job_id": job_id or f"enterprise_job_{uuid4().hex[:10]}", "organization_id": organization_id,
        "title": item.title, "domain": item.domain, "city": item.city, "headcount": item.headcount,
        "responsibilities": item.responsibilities, "required_skills": item.required_skills,
        "preferred_skills": item.preferred_skills, "experience_min_years": item.experience_min_years,
        "education": item.education, "salary_min_k": item.salary_min_k, "salary_max_k": item.salary_max_k,
        "status": item.status, "updated_at": now,
    }


def _enterprise_graph(job: dict) -> JobGraph:
    evidence = [SkillEvidence(source_id=job["job_id"], snippet=text, confidence=0.95) for text in (job.get("responsibilities") or [job["title"]])]
    skills = [
        SkillRequirement(skill_id=f"{job['job_id']}:{name}", name=name, category="core", required_level="intermediate", importance="required", evidence=evidence[:2], confidence=0.95)
        for name in job.get("required_skills", []) if is_valid_skill(name)
    ] + [
        SkillRequirement(skill_id=f"{job['job_id']}:{name}", name=name, category="plus", required_level="basic", importance="preferred", evidence=evidence[:1], confidence=0.82)
        for name in job.get("preferred_skills", []) if is_valid_skill(name)
    ]
    return JobGraph(job_id=job["job_id"], job_title=job["title"], domain=job["domain"], responsibilities=job.get("responsibilities") or [], skills=skills, scenarios=job.get("responsibilities") or [], evidence_summary={"source": "enterprise_job", "job_id": job["job_id"]})


def _market_graph(job) -> JobGraph:
    content = job.content or "\n".join(job.responsibilities) or job.title
    skills = extract_skills(content, job.skill_tags)

    def evidence_for(skill: str) -> list[SkillEvidence]:
        segments = [part.strip() for part in re.split(r"[\n。；;]", content) if part.strip()]
        matched = [part for part in segments if skill.casefold() in part.casefold()]
        snippets = (matched or segments)[:2]
        return [SkillEvidence(source_id=job.source_id, snippet=snippet[:260], confidence=max(0.5, job.quality_score / 100)) for snippet in snippets]

    return JobGraph(
        job_id=job.source_id,
        job_title=job.title,
        domain=job.domain,
        responsibilities=job.responsibilities,
        scenarios=job.responsibilities[:3],
        skills=[
            SkillRequirement(
                skill_id=f"{job.source_id}:{skill}", name=skill, category="market",
                required_level="intermediate", importance="required", evidence=evidence_for(skill),
                confidence=max(0.5, job.quality_score / 100),
            )
            for skill in skills if is_valid_skill(skill)
        ],
        evidence_summary={"source": "recruitment", "source_platform": job.source_platform, "source_url": job.source_url},
    )


def _report_result(report: MatchReport, *, source: str, job: dict) -> dict:
    job_snapshot = report.job_snapshot or {}
    evidence_summary = job_snapshot.get("evidence_summary") or {}
    return {
        "source": source,
        "job_id": report.target_job_id,
        "title": report.job_snapshot.get("job_title") or job.get("title"),
        "company": job.get("company") or "企业岗位库",
        "city": job.get("city") or "",
        "score": report.score,
        "level": report.level,
        "matched_skills": report.matched_skills,
        "missing_required_skills": report.missing_required_skills,
        "gaps": report.gaps,
        "strengths": report.strengths,
        "risks": report.risks,
        "interview_questions": report.follow_up_questions,
        "interview_recommendation": report.interview_recommendation,
        "executive_summary": report.executive_summary,
        "score_breakdown": report.score_breakdown,
        "evidence_coverage": report.evidence_coverage,
        "source_url": job.get("source_url"),
        "source_platform": job.get("source_platform") or evidence_summary.get("source_platform") or source,
        "required_skill_count": job_snapshot.get("required_skill_count", 0),
        "job_evidence": {
            "source_id": job.get("source_id") or report.target_job_id,
            "source_url": job.get("source_url"),
            "source_platform": job.get("source_platform") or evidence_summary.get("source_platform") or source,
            "evidence_summary": evidence_summary,
        },
    }


def _learning_job_key(title: str) -> str | None:
    normalized = title.casefold()
    for job_key, _, _, keywords in JOB_FAMILIES:
        if any(keyword.casefold() in normalized for keyword in keywords):
            return job_key
    return None


def _profile_from_parse(candidate_id: str, parsed: dict, resume_text: str) -> CandidateProfile:
    return CandidateProfile(
        candidate_id=candidate_id, target_job_id="", name=str(parsed.get("candidate_name") or "候选人"),
        skills=[item for item in parsed.get("skills", []) if isinstance(item, dict)], projects=parsed.get("projects") or [],
        evidence=parsed.get("evidence") or [{"type": "resume", "content": resume_text[:500], "confidence": 0.7}],
        confidence=0.75, summary=str(parsed.get("summary") or "简历结构化解析完成"),
        experience_years=parsed.get("experience_years"), education=parsed.get("education"),
        preferred_locations=parsed.get("preferred_locations") or [], expected_salary_k=parsed.get("expected_salary_k") or [],
        provider_status=parsed.get("provider_status") or {},
        resume_analysis={
            "job_intent": parsed.get("job_intent"),
            "work_history": parsed.get("work_history") or [],
            "certifications": parsed.get("certifications") or [],
            "quality_summary": parsed.get("quality_summary") or {},
            "semantic_profile": parsed.get("semantic_profile") or {},
        },
    )


def _resume_filename_display_name(filename: str | None) -> str | None:
    basename = str(filename or "").replace("\\", "/").rsplit("/", 1)[-1]
    stem = Path(basename).stem.strip()
    stem = re.sub(r"(?:[\s_-]*(?:个人简历|简历))+$", "", stem, flags=re.IGNORECASE)
    stem = re.sub(r"[\x00-\x1f]+", "", stem)
    stem = re.sub(r'[<>:"/\\|?*]+', "", stem)
    stem = re.sub(r"\s+", " ", stem).strip(" ._-")
    if not stem:
        return None
    normalized = re.sub(r"[\W_]+", "", stem, flags=re.UNICODE).casefold()
    generic_names = {
        "resume", "resumecv", "cv", "curriculumvitae", "candidate", "candidates", "applicant",
        "file", "document", "attachment", "upload", "untitled", "unknown", "unnamed", "newresume",
        "newdocument", "简历", "个人简历", "简历模板", "个人简历模板", "候选人", "求职者", "未命名", "附件",
    }
    if normalized in generic_names or normalized.isdigit():
        return None
    if re.fullmatch(
        r"(?:resume|cv|candidate|applicant|file|document|attachment|upload|untitled|unknown)(?:[-_ ]?(?:copy|final|new|v?\d+))?",
        stem,
        flags=re.IGNORECASE,
    ):
        return None
    return stem[:80]


def _apply_resume_filename_display_name(profile: CandidateProfile, parsed: dict, filename: str | None) -> None:
    if profile.name.strip() not in {"", "候选人"}:
        return
    display_name = _resume_filename_display_name(filename)
    if display_name:
        profile.name = display_name
        parsed["candidate_name"] = display_name


def _concise_resume_remote_error(error: ResumeParseError) -> str:
    reason = re.sub(r"\s+", " ", str(error)).strip()
    return reason[:240] or type(error).__name__


def _parse_resume_with_local_fallback(
    resume_text: str,
    *,
    candidate_id: str | None = None,
    source_name: str = "resume_text",
    images: list[dict[str, str]] | None = None,
) -> dict:
    parser = orchestrator.assessment_agent.resume_parser
    parse_kwargs = {
        "candidate_id": candidate_id,
        "source_name": source_name,
        "require_remote": True,
        "use_remote": True,
    }
    if images is not None:
        parse_kwargs["images"] = images
    try:
        return parser.parse(resume_text, **parse_kwargs)
    except ResumeParseError as exc:
        fallback_kwargs = {
            **parse_kwargs,
            "require_remote": False,
            "use_remote": False,
        }
        parsed = parser.parse(resume_text, **fallback_kwargs)
        provider_status = parsed.get("provider_status") if isinstance(parsed.get("provider_status"), dict) else {}
        parsed["provider_status"] = {
            **provider_status,
            "fallback": True,
            "mode": "local_evidence_fallback",
            "remote_error": _concise_resume_remote_error(exc),
        }
        return parsed


def ensure_domain(domain: str | None):
    if domain is not None and domain not in ALLOWED_DOMAINS:
        raise HTTPException(status_code=422, detail=f"domain must be one of {sorted(ALLOWED_DOMAINS)}")


@app.get("/api/health")
def health():
    return {"status": "ok", "service": "讯飞智聘"}


@app.get("/api/speech/status")
def speech_status(_user: dict = Depends(current_user)):
    return speech_service.status()


@app.websocket("/api/speech/stream")
async def speech_stream(websocket: WebSocket):
    await websocket.accept()
    try:
        start_message = await asyncio.wait_for(websocket.receive_json(), timeout=10)
        if start_message.get("type") != "start":
            await websocket.send_json({"type": "error", "message": "缺少语音识别启动信息"})
            await websocket.close(code=4400)
            return

        user = auth_service.resolve_token(str(start_message.get("token") or ""))
        if not user:
            await websocket.send_json({"type": "error", "message": "登录状态已失效，请重新登录"})
            await websocket.close(code=4401)
            return

        hotwords = start_message.get("hotwords")
        if not isinstance(hotwords, list):
            hotwords = []
        await websocket.send_json({"type": "status", "state": "loading", "message": "正在准备本地语音模型"})
        session = await asyncio.to_thread(speech_service.start_session, hotwords)
        await websocket.send_json({"type": "ready", **speech_service.status()})

        while True:
            event = await websocket.receive()
            if event["type"] == "websocket.disconnect":
                break
            pcm_bytes = event.get("bytes")
            if pcm_bytes is not None:
                if len(pcm_bytes) > 64 * 1024:
                    await websocket.send_json({"type": "error", "message": "单次音频数据过大"})
                    await websocket.close(code=4400)
                    return
                transcript = await asyncio.to_thread(session.feed_pcm16, pcm_bytes)
                if transcript is not None:
                    await websocket.send_json(transcript.as_payload())
                continue

            text = event.get("text")
            if not text:
                continue
            try:
                command = json.loads(text)
            except json.JSONDecodeError:
                command = {}
            if command.get("type") == "stop":
                transcript = await asyncio.to_thread(session.finalize)
                await websocket.send_json(transcript.as_payload())
                await websocket.send_json({"type": "complete", "text": transcript.text})
                await websocket.close(code=1000)
                return
    except TimeoutError:
        await websocket.send_json({"type": "error", "message": "语音识别连接超时，请重试"})
        await websocket.close(code=4408)
    except WebSocketDisconnect:
        return
    except SpeechModelUnavailable as exc:
        await websocket.send_json({"type": "error", "message": str(exc)})
        await websocket.close(code=1011)
    except RuntimeError:
        return


@app.post("/api/auth/login")
def login(payload: AuthLoginRequest):
    try:
        return auth_service.login(
            email=payload.email,
            password=payload.password,
            role=payload.role,
        )
    except AuthError as exc:
        raise HTTPException(status_code=401, detail=str(exc)) from exc


@app.post("/api/auth/register", status_code=201)
def register(payload: AuthRegisterRequest):
    try:
        user = auth_service.register(
            email=payload.email,
            password=payload.password,
            role=payload.role,
            display_name=payload.display_name,
            organization_name=payload.organization_name,
        )
        return auth_service.login(
            email=payload.email,
            password=payload.password,
            role=user["role"],
        )
    except AuthError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@app.get("/api/auth/me")
def auth_me(user: dict = Depends(current_user)):
    return user


@app.post("/api/auth/logout", status_code=204)
def logout(
    authorization: str | None = Header(default=None),
    _user: dict = Depends(current_user),
):
    auth_service.logout(_bearer_token(authorization))
    return None


def _decision_job_family(title: str | None, domain: str | None = None) -> str:
    text = f"{title or ''} {domain or ''}".casefold()
    families = (
        (("agent", "智能体"), "AI Agent工程师"),
        (("大模型", "llm", "rag"), "大模型应用工程师"),
        (("算法", "机器学习", "深度学习", "nlp", "视觉"), "AI算法工程师"),
        (("数据仓库", "数仓"), "数据仓库工程师"),
        (("大数据", "数据开发", "spark", "flink"), "大数据开发工程师"),
        (("数据分析", "商业分析", "bi"), "数据分析师"),
        (("java", "后端", "微服务"), "后端开发工程师"),
        (("云原生", "devops", "运维", "sre"), "云原生与DevOps工程师"),
        (("测试开发", "自动化测试", "测试工程"), "测试开发工程师"),
        (("物联网", "iot"), "物联网工程师"),
        (("嵌入式", "单片机", "stm32"), "嵌入式工程师"),
        (("智能系统", "智能网联"), "智能系统工程师"),
    )
    for keywords, family in families:
        if any(keyword in text for keyword in keywords):
            return family
    cleaned = re.sub(r"[\[【（(].*?[\]】）)]", "", title or "").strip()
    cleaned = re.sub(r"(高级|资深|中级|初级|实习|专家|负责人|经理)", "", cleaned).strip()
    return cleaned[:24] or domain or "其他数字技术岗位"


def _market_decision_summary(jobs: list, versions: list[dict] | None = None) -> dict:
    active_jobs = [job for job in jobs if job.job_status == "active"]
    versions = versions if versions is not None else store.list_all_job_versions()
    changed_versions = [item for item in versions if int(item.get("version_number") or 1) > 1]
    current_skill_counts = Counter(
        skill for job in active_jobs for skill in job.skill_tags if is_valid_skill(skill)
    )
    baseline_skill_counts = Counter()
    added_skill_counts = Counter()
    removed_skill_counts = Counter()
    for version in versions:
        snapshot = version.get("snapshot") or {}
        if int(version.get("version_number") or 1) == 1:
            baseline_skill_counts.update(
                skill for skill in snapshot.get("skill_tags") or [] if is_valid_skill(skill)
            )
        skill_change = (version.get("changes") or {}).get("skill_tags") or {}
        before = {skill for skill in skill_change.get("before") or [] if is_valid_skill(skill)}
        after = {skill for skill in skill_change.get("after") or [] if is_valid_skill(skill)}
        added_skill_counts.update(after - before)
        removed_skill_counts.update(before - after)

    current_family_counts = Counter(_decision_job_family(job.title, job.domain) for job in active_jobs)
    new_family_counts = Counter(
        _decision_job_family(job.title, job.domain)
        for job in jobs if job.change_type == "new" and job.job_status == "active"
    )
    offline_family_counts = Counter(
        _decision_job_family(job.title, job.domain)
        for job in jobs if job.job_status != "active"
    )

    def job_items(counter: Counter, direction: str) -> list[dict]:
        return [
            {
                "name": name,
                "signal_count": signal_count,
                "current_count": current_family_counts.get(name, 0),
                "direction": direction,
                "reason": f"监测到 {signal_count} 条{'新增在招' if direction == 'up' else '下线或退出'}证据",
            }
            for name, signal_count in counter.most_common(5) if signal_count > 0
        ]

    def skill_items(kind: str) -> list[dict]:
        selected = []
        names = set(current_skill_counts) | set(added_skill_counts) | set(removed_skill_counts)
        for name in names:
            current = current_skill_counts[name]
            baseline = baseline_skill_counts[name]
            added = added_skill_counts[name]
            removed = removed_skill_counts[name]
            matches = {
                "new": added > 0 and baseline == 0 and current > 0,
                "enhanced": baseline > 0 and added > removed,
                "weakened": current > 0 and removed > added,
                "exited": removed > 0 and current == 0,
            }
            if not matches[kind]:
                continue
            selected.append({
                "name": name,
                "delta": added - removed,
                "current_count": current,
                "coverage": round(current * 100 / max(len(active_jobs), 1), 1),
                "added_count": added,
                "removed_count": removed,
            })
        selected.sort(key=lambda item: (abs(item["delta"]), item["current_count"]), reverse=True)
        return selected[:5]

    common_abilities = [
        {
            "name": name,
            "job_count": count,
            "coverage": round(count * 100 / max(len(active_jobs), 1), 1),
            "stage": "普遍要求" if count / max(len(active_jobs), 1) >= 0.1 else "形成中",
        }
        for name, count in current_skill_counts.most_common(6)
    ]
    source_counts = Counter(job.source_platform for job in jobs)
    return {
        "growth_jobs": job_items(new_family_counts, "up"),
        "decline_jobs": job_items(offline_family_counts, "down"),
        "skill_changes": {kind: skill_items(kind) for kind in ("new", "enhanced", "weakened", "exited")},
        "common_abilities": common_abilities,
        "evidence": {
            "active_jobs": len(active_jobs),
            "version_snapshots": len(versions),
            "changed_jobs": len({item.get("job_id") for item in changed_versions}),
            "primary_sources": sum(source_counts.get(source, 0) > 0 for source in RECRUITMENT_SOURCE_IDS),
            "average_quality": round(sum(job.quality_score for job in jobs) / max(len(jobs), 1)),
        },
        "method": "岗位趋势基于新增、下线与版本变化信号；技能趋势基于同一岗位版本前后技能证据差异。",
        "generated_at": datetime.now(UTC).isoformat(),
    }


def _enterprise_standard_alignment(user: dict, jobs: list) -> dict:
    organization_id = user.get("organization_id") or user["user_id"]
    enterprise_jobs = [
        item for item in store._enterprise_rows("enterprise_jobs", organization_id)
        if item.get("status") == "active"
    ]
    if not enterprise_jobs:
        return {
            "status": "not_configured",
            "headline": "尚未建立企业岗位标准",
            "detail": "录入企业在招岗位后，系统将与招聘市场的当前能力要求逐项对照。",
            "total": 0,
            "outdated_count": 0,
            "items": [],
        }
    market_by_family: dict[str, Counter] = {}
    for job in jobs:
        if job.job_status == "active":
            family = _decision_job_family(job.title, job.domain)
            market_by_family.setdefault(family, Counter()).update(
                skill for skill in job.skill_tags if is_valid_skill(skill)
            )
    items = []
    for job in enterprise_jobs:
        family = _decision_job_family(job.get("title"), job.get("domain"))
        market_skills = [name for name, _ in market_by_family.get(family, Counter()).most_common(8)]
        own_skills = set(job.get("required_skills") or []) | set(job.get("preferred_skills") or [])
        gaps = [name for name in market_skills if name not in own_skills]
        items.append({
            "job_id": job.get("job_id"),
            "title": job.get("title"),
            "status": "outdated" if len(gaps) >= 2 else "current",
            "missing_market_skills": gaps[:4],
        })
    outdated_count = sum(item["status"] == "outdated" for item in items)
    return {
        "status": "outdated" if outdated_count else "current",
        "headline": f"{outdated_count} 个岗位标准需要更新" if outdated_count else "企业岗位标准与市场基本同步",
        "detail": "按同岗位族高频能力要求与企业必备、加分技能进行对照。",
        "total": len(items),
        "outdated_count": outdated_count,
        "items": items[:5],
    }


@app.get("/api/portal/summary")
def portal_summary(user: dict = Depends(current_user)):
    jobs = list(store.jobs.values())
    active_jobs = [job for job in jobs if job.job_status == "active"]
    skills = {skill for job in active_jobs for skill in job.skill_tags if is_valid_skill(skill)}
    versions = store.list_all_job_versions()
    changed_jobs = {item["job_id"] for item in versions if item.get("version_number", 1) > 1}
    graph_stats = orchestrator.knowledge_graph_service.build_hierarchy_level(jobs)["stats"]
    common_metrics = {
        "jobs": len(jobs),
        "active_jobs": len(active_jobs),
        "skills": len(skills),
        "job_families": graph_stats.get("job_families", 0),
        "emerging_jobs": len(store.emerging_jobs),
        "version_snapshots": len(versions),
        "changed_jobs": len(changed_jobs),
        "community_signals": len(store.community_signals),
    }
    decision_summary = _market_decision_summary(jobs, versions)
    if user["role"] == "enterprise":
        return {
            "role": "enterprise",
            "organization": user.get("organization_name"),
            "metrics": {
                **common_metrics,
                "pending_audits": sum(
                    record.review_status == "pending"
                    for record in store.data_audits.values()
                ),
                "assessment_sessions": len(store.sessions),
            },
            "permissions": [
                "graph:read",
                "graph:review",
                "jobs:manage",
                "evolution:review",
                "candidates:assess",
                "reports:enterprise",
                "data:govern",
            ],
            "decision_summary": decision_summary,
            "standard_alignment": _enterprise_standard_alignment(user, jobs),
        }
    applications = auth_service.list_applications(user)
    return {
        "role": "seeker",
        "display_name": user["display_name"],
        "metrics": {
            "applications": len(applications),
            "active_applications": sum(
                item["status"] not in {"rejected", "withdrawn"}
                for item in applications
            ),
            "available_jobs": len(store.jobs),
            **common_metrics,
            "assessment_sessions": 0,
        },
        "permissions": [
            "graph:read",
            "evolution:read",
            "emerging_jobs:read",
            "resume:own",
            "assessment:self",
            "reports:self",
            "learning:self",
            "applications:self",
        ],
        "decision_summary": decision_summary,
    }


@app.get("/api/platform/summary")
def platform_summary():
    jobs = list(store.jobs.values())
    active_jobs = [job for job in jobs if job.job_status == "active"]
    skills = {skill for job in active_jobs for skill in job.skill_tags if is_valid_skill(skill)}
    skill_counts = Counter(skill for job in active_jobs for skill in job.skill_tags if is_valid_skill(skill))
    versions = store.list_all_job_versions()
    changed = [item for item in versions if item.get("version_number", 1) > 1]
    graph_stats = orchestrator.knowledge_graph_service.build_hierarchy_level(jobs)["stats"]
    return {
        "jobs": len(jobs),
        "active_jobs": len(active_jobs),
        "skills": len(skills),
        "job_families": graph_stats.get("job_families", 0),
        "emerging_jobs": len(store.emerging_jobs),
        "version_snapshots": len(versions),
        "changed_jobs": len({item["job_id"] for item in changed}),
        "change_events": len(changed),
        "community_signals": len(store.community_signals),
        "top_skills": [{"name": name, "job_count": count} for name, count in skill_counts.most_common(5)],
        "recent_job_changes": [
            {
                "source_id": job.source_id,
                "title": job.title,
                "domain": job.domain,
                "company": job.company,
                "city": job.city,
                "skill_tags": job.skill_tags[:5],
                "changed_at": job.changed_at,
                "change_type": job.change_type,
                "job_status": job.job_status,
                "source_url": job.source_url,
            }
            for job in sorted(
                (job for job in jobs if job.change_type != "unchanged"),
                key=lambda job: job.changed_at or job.collected_at,
                reverse=True,
            )[:5]
        ],
        "average_quality": round(sum(job.quality_score for job in jobs) / max(len(jobs), 1)),
        "sources": {
            source: sum(job.source_platform == source for job in jobs)
            for source in RECRUITMENT_SOURCE_IDS
        },
        "decision_summary": _market_decision_summary(jobs, versions),
    }


@app.get("/api/applications")
def list_applications(user: dict = Depends(current_user)):
    items = auth_service.list_applications(user)
    return {"total": len(items), "items": items}


@app.get("/api/assessment-records")
def list_assessment_records(user: dict = Depends(current_user)):
    items = store.list_assessment_records(user)
    return {"total": len(items), "items": items}


@app.put("/api/assessment-records/{record_id}")
def save_assessment_record(record_id: str, payload: AssessmentRecordRequest, user: dict = Depends(current_user)):
    record = {**payload.record, "id": record_id}
    try:
        return store.save_assessment_record(record, user)
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc


@app.delete("/api/assessment-records/{record_id}", status_code=204)
def delete_assessment_record(record_id: str, user: dict = Depends(current_user)):
    try:
        deleted = store.delete_assessment_record(record_id, user)
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    if not deleted:
        raise HTTPException(status_code=404, detail="assessment record not found")
    return None


@app.post("/api/applications", status_code=201)
def create_application(
    payload: ApplicationCreateRequest,
    user: dict = Depends(require_role("seeker")),
):
    try:
        return auth_service.create_application(
            user=user,
            job_id=payload.job_id,
            job_title=payload.job_title,
            company=payload.company,
            source_url=payload.source_url,
            match_score=payload.match_score,
            organization_id=payload.organization_id,
            snapshot=payload.snapshot,
        )
    except AuthError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


@app.patch("/api/applications/{application_id}")
def update_application(
    application_id: str,
    payload: ApplicationStatusRequest,
    user: dict = Depends(current_user),
):
    try:
        return auth_service.update_application(application_id, payload.status, user)
    except AuthError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc


@app.get("/api/llm/status")
def llm_status():
    return orchestrator.assessment_agent.llm.status()


@app.post("/api/data/import")
def import_data(payload: ImportRequest):
    batch = orchestrator.import_data(payload)
    store.index_batch(batch)
    return batch


@app.get("/api/data/sources")
def data_sources():
    return {
        "items": orchestrator.data_collection_agent.catalog(),
        "scope": ["人工智能", "大数据", "智能系统", "物联网"],
        "evidence_policy": {"recruitment": "primary", "community_signal": "auxiliary"},
    }


@app.get("/api/data/browser-sessions")
def browser_session_statuses():
    return {"items": browser_sessions.statuses()}


@app.post("/api/data/browser-sessions/{source_id}/start")
def start_browser_session(source_id: str):
    try:
        return browser_sessions.start(source_id)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except RuntimeError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc


@app.get("/api/data/cookie-sessions")
def cookie_session_statuses():
    return {"items": cookie_sessions.statuses()}


@app.put("/api/data/cookie-sessions/{source_id}")
def save_cookie_session(source_id: str, payload: CookieSessionRequest):
    try:
        return cookie_sessions.save(source_id, payload.cookies)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/api/data/cookie-sessions/{source_id}")
def clear_cookie_session(source_id: str):
    try:
        return cookie_sessions.clear(source_id)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.post("/api/data/collect", status_code=202)
def collect_data(payload: DataCollectionRequest, background_tasks: BackgroundTasks):
    run = orchestrator.data_collection_agent.prepare_run(payload)
    background_tasks.add_task(orchestrator.data_collection_agent.collect, payload, run=run)
    return run


@app.get("/api/data/collection-runs")
def collection_runs(limit: int = Query(default=20, ge=1, le=100)):
    runs = sorted(store.collection_runs.values(), key=lambda run: run.started_at, reverse=True)
    return {"total": len(runs), "items": runs[:limit]}


@app.get("/api/data/collection-runs/{run_id}")
def collection_run(run_id: str):
    run = store.collection_runs.get(run_id)
    if run is None:
        raise HTTPException(status_code=404, detail="collection run not found")
    return run


@app.get("/api/data/full-crawl-status")
def full_crawl_status():
    path = Path("data/full_recruitment_crawl_state.json")
    if not path.exists():
        return {"status": "not_started", "completed_count": 0, "total_count": 0}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {"status": "unavailable", "completed_count": 0, "total_count": 0}
    mode = value.get("mode", "discovery")
    stage_count = 2 if mode == "both" else 1
    source_configs = value.get("original_strategy", {}).get("sources", {})
    calculated_total = sum(
        max(1, len(source_configs.get(source_id, {}).get("cities") or ["全国"]))
        * len(source_configs.get(source_id, {}).get("job_queries") or [])
        for source_id in ("boss", "zhaopin")
    ) * stage_count
    total = int(value.get("total_count") or calculated_total)
    completed = len(value.get("completed", []))
    return {
        "status": value.get("status", "unknown"),
        "mode": mode,
        "stage": value.get("stage"),
        "source_id": value.get("source_id"),
        "city": value.get("city"),
        "query": value.get("query"),
        "completed_count": completed,
        "total_count": total,
        "remaining_count": max(0, total - completed),
        "progress_percent": round(completed * 100 / total, 2) if total else 0,
        "evidence_count": len(value.get("completion_evidence", {})),
        "totals": value.get("totals", {}),
        "last_run": value.get("last_run"),
        "last_error": value.get("last_error"),
        "retry_in_seconds": value.get("retry_in_seconds"),
        "next_cycle_at": value.get("next_cycle_at"),
        "cycle_type": value.get("cycle_type", "initial_discovery"),
        "updated_at": value.get("updated_at"),
    }


@app.get("/api/data/community-signals")
def community_signals(
    domain: str | None = None,
    source_platform: str | None = None,
    limit: int = Query(default=50, ge=1, le=200),
):
    ensure_domain(domain)
    signals = store.list_community_signals(domain=domain, source_platform=source_platform)
    return {"total": len(signals), "items": signals[:limit], "evidence_role": "auxiliary"}


@app.get("/api/data/schedule")
def collection_schedule():
    return collection_scheduler.status()


@app.put("/api/data/schedule")
def update_collection_schedule(payload: CollectionScheduleRequest):
    return collection_scheduler.configure(payload)


@app.get("/api/data/strategy")
def collection_strategy():
    return collection_scheduler.strategy_status()


@app.put("/api/data/strategy")
def update_collection_strategy(payload: CollectionStrategyRequest):
    return collection_scheduler.configure_strategy(payload)


@app.post("/api/data/revalidate")
def revalidate_data(
    min_quality: float = Query(default=65, ge=0, le=100),
    min_signal_quality: float = Query(default=55, ge=0, le=100),
):
    return orchestrator.data_collection_agent.revalidate_existing_jobs(min_quality, min_signal_quality)


@app.post("/api/data/evidence/rebuild")
def rebuild_data_evidence():
    return orchestrator.data_collection_agent.rebuild_evidence_intelligence()


@app.get("/api/data/quality-rejections")
def quality_rejections(limit: int = Query(default=100, ge=1, le=500)):
    records = sorted(store.quality_rejections.values(), key=lambda item: item.rejected_at, reverse=True)
    status_summary = {
        "pending": sum(item.review_status == "pending" for item in records),
        "approved": sum(item.review_status == "approved" for item in records),
        "rejected": sum(item.review_status == "rejected" for item in records),
    }
    return {"total": len(records), "status_summary": status_summary, "items": records[:limit]}


@app.get("/api/data/audits")
def data_audits(
    lane: str | None = None,
    record_type: str | None = None,
    status: str | None = None,
    source_platform: str | None = None,
    query: str | None = None,
    limit: int = Query(default=200, ge=1, le=1000),
    offset: int = Query(default=0, ge=0),
):
    records = sorted(store.data_audits.values(), key=lambda item: item.updated_at, reverse=True)
    lane_filters = {
        "pending": lambda record: (
            record.status in {"manual_review", "pending_semantic_review"}
            and record.record_type not in {"copy_cluster", "job_change", "change_review"}
        ),
        "rejected": lambda record: record.record_type == "quality_rejection" or record.status in {"rejected", "quarantined"},
        "corrections": lambda record: record.record_type == "field_correction",
        "copies": lambda record: record.record_type == "copy_cluster",
        "changes": lambda record: record.record_type in {"job_change", "change_review"},
    }
    if lane and lane not in lane_filters:
        raise HTTPException(status_code=422, detail="invalid audit lane")
    all_records = list(store.data_audits.values())
    lane_summary = {
        lane_name: sum(predicate(record) for record in all_records)
        for lane_name, predicate in lane_filters.items()
    }
    if lane:
        records = [record for record in records if lane_filters[lane](record)]
    if record_type:
        records = [record for record in records if record.record_type == record_type]
    if status:
        records = [record for record in records if record.status == status or record.review_status == status]
    if source_platform:
        records = [record for record in records if record.source_platform == source_platform]
    if query:
        needle = query.casefold()
        records = [
            record for record in records
            if needle in f"{record.title} {record.source_id} {' '.join(record.reasons)}".casefold()
        ]
    type_summary = {}
    status_summary = {}
    for record in all_records:
        type_summary[record.record_type] = type_summary.get(record.record_type, 0) + 1
        status_summary[record.status] = status_summary.get(record.status, 0) + 1
    return {
        "total": len(records),
        "offset": offset,
        "limit": limit,
        "type_summary": type_summary,
        "status_summary": status_summary,
        "lane_summary": lane_summary,
        "items": records[offset: offset + limit],
    }


@app.get("/api/data/audits/{audit_id}")
def data_audit(audit_id: str):
    record = store.data_audits.get(audit_id)
    if record is None:
        raise HTTPException(status_code=404, detail="data audit record not found")
    return record


@app.post("/api/data/audits/{audit_id}/review")
def review_data_audit(
    audit_id: str,
    payload: DataReviewRequest,
    background_tasks: BackgroundTasks,
):
    if payload.status not in {"approved", "rejected"}:
        raise HTTPException(status_code=422, detail="status must be approved or rejected")
    record = store.review_data_audit(
        audit_id, payload.status, payload.note, payload.decision, payload.corrections,
    )
    if record is None:
        raise HTTPException(status_code=404, detail="data audit record not found")
    background_tasks.add_task(_reprocess_governance_record, record)
    return {"record": record, "processing": "queued"}


def _reprocess_governance_record(record) -> None:
    orchestrator.data_collection_agent.reprocess_data_audit(record)
    if record.review_decision == "new_job_candidate" and record.review_status == "approved":
        orchestrator.discover_emerging_jobs(
            EmergingJobDiscoveryRequest(minimum_source_count=1, minimum_confidence=0.35)
        )


@app.get("/api/data/copy-clusters")
def copy_clusters(limit: int = Query(default=100, ge=1, le=500)):
    cached = store.get_setting("copy_clusters", {})
    items = cached.get("items") or orchestrator.data_collection_agent.evidence_service.copy_clusters(store.jobs.values())
    return {"total": len(items), "items": items[:limit], "updated_at": cached.get("updated_at")}


@app.get("/api/signals/technology-trends")
def technology_trends(
    domain: str | None = None,
    status: str | None = None,
    limit: int = Query(default=80, ge=1, le=300),
):
    ensure_domain(domain)
    cached = store.get_setting("technology_trends", {})
    items = cached.get("items") or orchestrator.data_collection_agent.evidence_service.technology_trends(
        store.jobs.values(), store.community_signals.values(),
    )
    if domain:
        items = [item for item in items if item["domain"] == domain]
    if status:
        items = [item for item in items if item["status"] == status]
    return {
        "total": len(items),
        "items": items[:limit],
        "updated_at": cached.get("updated_at"),
        "evidence_policy": {"recruitment": "fact", "community_signal": "auxiliary_only"},
    }


@app.get("/api/collection/coverage")
def collection_coverage():
    sources = []
    for source_id, state_file in SOURCE_STATE_PATHS.items():
        strategy = orchestrator.data_collection_agent.source_strategies.get(source_id, {})
        total_combinations = max(1, len(strategy.get("cities") or ["全国"])) * max(
            1, len(strategy.get("job_queries") or [])
        )
        try:
            state = json.loads(state_file.read_text(encoding="utf-8")) if state_file.exists() else {}
        except (OSError, json.JSONDecodeError):
            state = {}
        visited = len(state.get("page_cursors", {}))
        sources.append({
            "source_id": source_id,
            "total_combinations": total_combinations,
            "visited_combinations": min(total_combinations, visited),
            "coverage_percent": round(min(total_combinations, visited) * 100 / total_combinations, 2),
            "seen_jobs": len(state.get("seen", {})),
            "cycle_started_at": state.get("cycle_started_at"),
            "mode": strategy.get("mode", "hybrid"),
            "scope_statement": "仅表示已配置城市与搜索词组合覆盖，不代表招聘网站全站覆盖",
        })
    return {"sources": sources, "generated_at": datetime.now(UTC)}


@app.patch("/api/data/quality-rejections/{rejection_id}")
def review_quality_rejection(rejection_id: str, payload: DataReviewRequest, background_tasks: BackgroundTasks):
    if payload.status not in {"approved", "rejected"}:
        raise HTTPException(status_code=422, detail="status must be approved or rejected")
    record = store.review_quality_rejection(
        rejection_id, payload.status, payload.note, payload.decision, payload.corrections,
    )
    if record is None:
        raise HTTPException(status_code=404, detail="quality audit record not found")
    background_tasks.add_task(orchestrator.data_collection_agent.reprocess_reviewed_record, record)
    return {"record": record, "processing": "queued"}


@app.get("/api/data/audit-learning")
def audit_learning_status():
    return store.audit_learning_status()


@app.get("/api/jobs")
def list_jobs(
    city: str | None = None,
    domain: str | None = None,
    keyword: str | None = None,
    source_platform: str | None = None,
    status: str | None = None,
    min_quality: float = Query(default=0, ge=0, le=100),
    changed_within_days: int | None = Query(default=None, ge=1, le=3650),
    active_only: bool = True,
    limit: int = Query(default=50, ge=1, le=200),
    offset: int = Query(default=0, ge=0),
):
    ensure_domain(domain)
    if status is not None and status not in {"active", "offline"}:
        raise HTTPException(status_code=422, detail="status must be active or offline")
    jobs = store.list_jobs(keyword=keyword, active_only=active_only)
    if domain:
        jobs = [job for job in jobs if job.domain == domain]
    if source_platform:
        jobs = [job for job in jobs if job.source_platform == source_platform]
    if city:
        jobs = [job for job in jobs if job.city == city]
    if status:
        jobs = [job for job in jobs if job.job_status == status]
    if min_quality:
        jobs = [job for job in jobs if job.quality_score >= min_quality]
    if changed_within_days:
        cutoff = datetime.now(UTC).timestamp() - changed_within_days * 86400
        jobs = [job for job in jobs if job.changed_at and job.changed_at.timestamp() >= cutoff]
    return {"total": len(jobs), "offset": offset, "limit": limit, "items": jobs[offset : offset + limit]}


@app.get("/api/enterprise-jobs")
def list_enterprise_jobs(user: dict = Depends(require_role("enterprise"))):
    organization_id = user.get("organization_id") or user["user_id"]
    items = store._enterprise_rows("enterprise_jobs", organization_id)
    return {"total": len(items), "items": items}


@app.post("/api/enterprise-jobs", status_code=201)
def create_enterprise_job(payload: EnterpriseJobRequest, user: dict = Depends(require_role("enterprise"))):
    ensure_domain(payload.domain)
    if payload.status not in {"active", "paused", "closed"}:
        raise HTTPException(status_code=422, detail="status must be active, paused or closed")
    organization_id = user.get("organization_id") or user["user_id"]
    item = _enterprise_job_payload(payload, organization_id)
    return store.save_enterprise_record("enterprise_jobs", item["job_id"], organization_id, item)


@app.put("/api/enterprise-jobs/{job_id}")
def update_enterprise_job(job_id: str, payload: EnterpriseJobRequest, user: dict = Depends(require_role("enterprise"))):
    ensure_domain(payload.domain)
    if payload.status not in {"active", "paused", "closed"}:
        raise HTTPException(status_code=422, detail="status must be active, paused or closed")
    organization_id = user.get("organization_id") or user["user_id"]
    existing = store._enterprise_rows("enterprise_jobs", organization_id, item_id=job_id)
    if not existing:
        raise HTTPException(status_code=404, detail="enterprise job not found")
    item = _enterprise_job_payload(payload, organization_id, job_id=job_id)
    return store.save_enterprise_record("enterprise_jobs", job_id, organization_id, item)


@app.get("/api/enterprise-jobs/template")
def enterprise_job_template():
    book = Workbook()
    sheet = book.active
    sheet.title = "招聘岗位"
    headers = ["岗位名称", "所属领域", "工作城市", "招聘人数", "岗位职责（每行一条）", "必备技能（逗号分隔）", "加分技能（逗号分隔）", "最低经验年限", "学历", "最低月薪K", "最高月薪K", "状态"]
    sheet.append(headers)
    sheet.append(["AI Agent 工程师", "人工智能", "杭州", 2, "负责智能体应用开发", "Python,LangGraph,RAG", "FastAPI,向量数据库", 3, "本科", 25, 45, "active"])
    output = BytesIO()
    book.save(output)
    return StreamingResponse(BytesIO(output.getvalue()), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": "attachment; filename=enterprise-job-template.xlsx"})


@app.post("/api/enterprise-jobs/import")
async def import_enterprise_jobs(file: UploadFile = File(...), user: dict = Depends(require_role("enterprise"))):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=422, detail="请上传 .xlsx 岗位模板")
    try:
        book = load_workbook(BytesIO(await file.read()), read_only=True, data_only=True)
        rows = list(book.active.iter_rows(values_only=True))
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"无法读取 Excel 文件：{exc}") from exc
    if not rows:
        raise HTTPException(status_code=422, detail="Excel 文件为空")
    headers = {str(value).strip(): index for index, value in enumerate(rows[0]) if value}
    required_headers = {"岗位名称", "所属领域", "工作城市", "必备技能（逗号分隔）"}
    if not required_headers.issubset(headers):
        raise HTTPException(status_code=422, detail="模板缺少必填列")
    organization_id = user.get("organization_id") or user["user_id"]
    created, errors = [], []
    def value(row, name, default=""):
        index = headers.get(name)
        return row[index] if index is not None and index < len(row) and row[index] is not None else default
    for number, row in enumerate(rows[1:], start=2):
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        try:
            split = lambda raw: [part.strip() for part in re.split(r"[,，;；]", str(raw or "")) if part.strip()]
            raw = EnterpriseJobRequest(title=str(value(row, "岗位名称")).strip(), domain=str(value(row, "所属领域")).strip(), city=str(value(row, "工作城市")).strip(), headcount=int(value(row, "招聘人数", 1) or 1), responsibilities=[part.strip() for part in str(value(row, "岗位职责（每行一条）", "")).splitlines() if part.strip()], required_skills=split(value(row, "必备技能（逗号分隔）")), preferred_skills=split(value(row, "加分技能（逗号分隔）")), experience_min_years=int(value(row, "最低经验年限")) if value(row, "最低经验年限") != "" else None, education=str(value(row, "学历", "")).strip() or None, salary_min_k=int(value(row, "最低月薪K")) if value(row, "最低月薪K") != "" else None, salary_max_k=int(value(row, "最高月薪K")) if value(row, "最高月薪K") != "" else None, status=str(value(row, "状态", "active")).strip() or "active")
            ensure_domain(raw.domain)
            if raw.status not in {"active", "paused", "closed"}:
                raise ValueError("状态必须为 active、paused 或 closed")
            item = _enterprise_job_payload(raw, organization_id)
            store.save_enterprise_record("enterprise_jobs", item["job_id"], organization_id, item)
            created.append(item)
        except Exception as exc:
            errors.append({"row": number, "message": str(exc)})
    return {"created": len(created), "items": created, "errors": errors}


def _parse_target_job_ids(target_job_id: str | None, target_job_ids: str | None) -> list[str]:
    selected: list[str] = []
    if target_job_ids:
        try:
            raw_ids = json.loads(target_job_ids)
        except json.JSONDecodeError as exc:
            raise HTTPException(status_code=422, detail="target_job_ids 必须是 JSON 数组") from exc
        if not isinstance(raw_ids, list):
            raise HTTPException(status_code=422, detail="target_job_ids 必须是 JSON 数组")
        selected.extend(str(item).strip() for item in raw_ids if str(item).strip())
    if target_job_id and target_job_id.strip():
        selected.insert(0, target_job_id.strip())
    selected = list(dict.fromkeys(selected))
    if not selected:
        raise HTTPException(status_code=422, detail="请至少选择一个投递岗位")
    if len(selected) > 10:
        raise HTTPException(status_code=422, detail="单份简历最多可同时投递 10 个岗位")
    return selected


def _recommend_roles(profile: CandidateProfile, target_jobs: list[dict]) -> dict:
    candidates: list[dict] = []
    reports_by_job: dict[str, dict] = {}
    for job in target_jobs:
        report = orchestrator.match_agent.run(_enterprise_graph(job), profile)
        reports_by_job[job["job_id"]] = report.model_dump(mode="json")
        candidates.append(_report_result(report, source="enterprise", job=job))
    ranked = sorted(candidates, key=lambda item: (-item["score"], len(item["missing_required_skills"]), item["title"]))
    primary_job_id = target_jobs[0]["job_id"]
    return {
        "target_report": reports_by_job[primary_job_id],
        "target_reports": reports_by_job,
        "all_job_results": ranked,
        "recommendations": [item for item in ranked if item["job_id"] != primary_job_id][:5],
        "best_match": ranked[0] if ranked else None,
    }


def _has_completed_candidate_assessment(candidate_id: str, profile: CandidateProfile, user: dict) -> bool:
    if profile.competency_dimensions or profile.behavior_signals:
        return True
    for record in store.list_assessment_records(user):
        if str(record.get("candidate_id") or "") != str(candidate_id):
            continue
        record_profile = record.get("profile") if isinstance(record.get("profile"), dict) else {}
        if (
            record_profile.get("match_report")
            or record_profile.get("competency_dimensions")
            or record_profile.get("behavior_signals")
        ):
            return True
    return False


@app.post("/api/recruitment/candidates/upload")
async def upload_recruitment_candidates(
    files: list[UploadFile] = File(...),
    target_job_id: str | None = Form(default=None),
    target_job_ids: str | None = Form(default=None),
    user: dict = Depends(require_role("enterprise")),
):
    organization_id = user.get("organization_id") or user["user_id"]
    selected_job_ids = _parse_target_job_ids(target_job_id, target_job_ids)
    active_jobs = {
        job["job_id"]: job
        for job in store._enterprise_rows("enterprise_jobs", organization_id)
        if job.get("status") == "active"
    }
    missing_job_ids = [job_id for job_id in selected_job_ids if job_id not in active_jobs]
    if missing_job_ids:
        raise HTTPException(status_code=404, detail=f"企业招聘岗位不存在或已停止招聘：{', '.join(missing_job_ids)}")
    target_jobs = [active_jobs[job_id] for job_id in selected_job_ids]
    items, errors = [], []
    for file in files:
        try:
            extracted = extract_resume_text(file.filename or "resume", await file.read())
            candidate_id = f"candidate_{uuid4().hex[:10]}"
            parsed = _parse_resume_with_local_fallback(
                str(extracted["text"]),
                candidate_id=candidate_id,
                source_name=str(extracted["filename"]),
            )
            profile = _profile_from_parse(candidate_id, parsed, str(extracted["text"]))
            _apply_resume_filename_display_name(profile, parsed, str(extracted["filename"]))
            store.save_candidate_profile(profile, user)
            result = _recommend_roles(profile, target_jobs)
            analysis = store.save_gap_analysis({
                "analysis_id": f"gap_{uuid4().hex[:12]}",
                "candidate_id": candidate_id,
                "scope": "enterprise_selected_jobs",
                "target_job_id": selected_job_ids[0],
                "target_job_ids": selected_job_ids,
                "evaluated_job_count": len(result["all_job_results"]),
                "data_policy": "real_recruitment_evidence_only",
                "analysis_method": "llm_resume_profile_plus_evidence_weighted_skill_matching",
                "source_counts": {"enterprise": len(result["all_job_results"])},
                "results": result["all_job_results"],
                "top_matches": result["all_job_results"][:5],
                "best_match": result["best_match"],
                "profile_report": profile.model_dump(mode="json"),
                "difference_report": result["all_job_results"],
            }, user)
            record = {"candidate_id": candidate_id, "name": profile.name, "target_job_id": selected_job_ids[0], "target_job_ids": selected_job_ids, "source_name": extracted["filename"], "resume_parse": parsed, "profile": profile.model_dump(mode="json"), "screening": result["target_report"], "target_screenings": result["target_reports"], "all_job_results": result["all_job_results"], "gap_analysis_id": analysis["analysis_id"], "recommendations": result["recommendations"], "interview_questions": [question for item in result["all_job_results"][:3] for question in item["interview_questions"]][:10], "profile_report": profile.model_dump(mode="json"), "difference_report": result["all_job_results"], "status": "pending_review", "created_at": datetime.now(UTC).isoformat()}
            store.save_enterprise_record("recruitment_candidates", candidate_id, organization_id, record)
            items.append(record)
        except Exception as exc:
            errors.append({"filename": file.filename, "message": str(exc)})
    return {"created": len(items), "items": items, "errors": errors}


@app.get("/api/recruitment/candidates")
def list_recruitment_candidates(target_job_id: str | None = None, user: dict = Depends(require_role("enterprise"))):
    organization_id = user.get("organization_id") or user["user_id"]
    items = store._enterprise_rows("recruitment_candidates", organization_id)
    if target_job_id:
        items = [
            item for item in items
            if target_job_id in (item.get("target_job_ids") or [item.get("target_job_id")])
        ]
    return {"total": len(items), "items": sorted(items, key=lambda item: item.get("screening", {}).get("score", 0), reverse=True)}


@app.post("/api/recruitment/candidates/{candidate_id}/feedback")
def save_recruitment_feedback(candidate_id: str, payload: RecruitmentFeedbackRequest, user: dict = Depends(require_role("enterprise"))):
    organization_id = user.get("organization_id") or user["user_id"]
    records = store._enterprise_rows("recruitment_candidates", organization_id, item_id=candidate_id)
    if not records:
        raise HTTPException(status_code=404, detail="候选人不存在")
    feedback = {"feedback_id": f"feedback_{uuid4().hex[:10]}", "candidate_id": candidate_id, "decision": payload.decision, "note": payload.note, "ratings": payload.ratings, "created_at": datetime.now(UTC).isoformat()}
    store.save_enterprise_record("recruitment_feedback", feedback["feedback_id"], organization_id, feedback)
    record = {**records[0], "status": payload.decision, "latest_feedback": feedback}
    store.save_enterprise_record("recruitment_candidates", candidate_id, organization_id, record)
    return record


@app.post("/api/gap-analysis")
def analyze_candidate_gaps(payload: GapAnalysisRequest, user: dict = Depends(current_user)):
    if not store.can_access_candidate(payload.candidate_id, user):
        raise HTTPException(status_code=403, detail="candidate profile access denied")
    profile = store.profiles.get(payload.candidate_id)
    if not profile:
        raise HTTPException(status_code=404, detail="candidate profile not found")
    provider_mode = str(profile.provider_status.get("mode") or "")
    has_completed_assessment = _has_completed_candidate_assessment(payload.candidate_id, profile, user)
    if provider_mode != "llm_resume_understanding" and not has_completed_assessment:
        raise HTTPException(status_code=409, detail="简历尚未完成大模型语义解析，不能生成正式岗位差距报告")

    results: list[dict] = []
    scope = "target_job"
    if user["role"] == "enterprise":
        organization_id = user.get("organization_id") or user["user_id"]
        jobs = [item for item in store._enterprise_rows("enterprise_jobs", organization_id) if item.get("status") == "active"]
        if not jobs:
            raise HTTPException(status_code=422, detail="企业当前没有在招岗位")
        for job in jobs:
            report = orchestrator.match_agent.run(_enterprise_graph(job), profile)
            results.append(_report_result(report, source="enterprise", job=job))
        scope = "enterprise_all_active"
    else:
        if not payload.analyze_all and not payload.target_job_id:
            raise HTTPException(status_code=422, detail="请选择目标岗位或启用全部岗位分析")
        market_jobs = [job for job in store.jobs.values() if job.job_status == "active"]
        if payload.target_job_id:
            market_jobs = [job for job in market_jobs if job.source_id == payload.target_job_id]
            if not market_jobs:
                raise HTTPException(status_code=404, detail="目标岗位不存在或已下线")
        for job in market_jobs:
            graph = _market_graph(job)
            if len(graph.skills) < 3:
                continue
            report = orchestrator.match_agent.run(graph, profile)
            results.append(_report_result(report, source="market", job=job.model_dump(mode="json")))
        scope = "market_all" if payload.analyze_all else "target_job"

    if not results:
        raise HTTPException(status_code=422, detail="没有找到具备至少 3 项真实 JD 技能证据的岗位，无法生成正式差距报告")

    ranked = sorted(results, key=lambda item: (-item["score"], len(item["missing_required_skills"]), item["title"]))
    visible_results = ranked if user["role"] == "enterprise" or not payload.analyze_all else ranked[:5]
    analysis = {
        "analysis_id": f"gap_{uuid4().hex[:12]}",
        "candidate_id": payload.candidate_id,
        "scope": scope,
        "target_job_id": payload.target_job_id,
        "evaluated_job_count": len(results),
        "data_policy": "real_recruitment_evidence_only",
        "analysis_method": "llm_resume_profile_plus_evidence_weighted_skill_matching",
        "source_counts": dict(Counter(item["source_platform"] for item in results)),
        "results": visible_results,
        "top_matches": ranked[:5],
        "best_match": ranked[0] if ranked else None,
        "profile_report": profile.model_dump(mode="json"),
        "difference_report": visible_results,
    }
    try:
        return store.save_gap_analysis(analysis, user)
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc


@app.get("/api/gap-analysis/latest/{candidate_id}")
def latest_candidate_gap_analysis(candidate_id: str, user: dict = Depends(current_user)):
    try:
        analysis = store.latest_gap_analysis(candidate_id, user)
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    if not analysis:
        raise HTTPException(status_code=404, detail="gap analysis not found")
    return analysis


def _hydrate_learning_course_covers(route: dict) -> dict:
    for step in route.get("steps", []):
        for course in step.get("courses", []):
            if course.get("cover_url"):
                continue
            stored = learning_service.store.get_course(str(course.get("course_id") or ""))
            if stored and stored.get("cover_url"):
                course["cover_url"] = stored["cover_url"]
    return route


@app.post("/api/learning/routes")
def create_learning_route(payload: LearningRouteRequest, user: dict = Depends(current_user)):
    try:
        analysis = store.get_gap_analysis(payload.analysis_id, user)
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    if not analysis:
        raise HTTPException(status_code=404, detail="gap analysis not found")
    requested_job_id = str(payload.job_id or "")
    selected = next(
        (item for item in analysis.get("results", []) if str(item.get("job_id")) == requested_job_id),
        None,
    ) if requested_job_id else analysis.get("best_match")
    if requested_job_id and not selected:
        raise HTTPException(status_code=422, detail=f"差距分析中不存在岗位 ID：{requested_job_id}")
    if not selected:
        raise HTTPException(status_code=422, detail="差距分析中没有可生成路线的岗位")
    weak_skills = [item.get("skill") for item in selected.get("gaps", []) if item.get("skill")][:8]
    if not weak_skills:
        raise HTTPException(status_code=422, detail="该岗位当前没有需要补齐的技能")

    if not learning_service.store.list_jobs():
        learning_service.sync_jobs()
    job_key = _learning_job_key(selected["title"])
    if not job_key or not learning_service.store.get_job(job_key):
        job_key = f"analysis:{selected['job_id']}"
        learning_service.store.replace_jobs([JobProfile(
            job_key=job_key,
            title=selected["title"],
            domain="岗位能力",
            job_count=1,
            skills=[JobSkill(name=skill, demand_count=1, demand_rate=100) for skill in weak_skills],
        )])
    route = learning_service.route(RouteRequest(
        job_key=job_key,
        weak_skills=weak_skills,
        current_level=payload.current_level,
        courses_per_skill=payload.courses_per_skill,
    )).model_dump(mode="json")
    route.update({"analysis_id": payload.analysis_id, "candidate_id": analysis["candidate_id"], "job_id": selected["job_id"]})
    return store.save_learning_route(_hydrate_learning_course_covers(route), payload.analysis_id, analysis["candidate_id"], user)


@app.get("/api/learning/courses")
def learning_courses(skill: str = Query(min_length=1), limit: int = Query(default=10, ge=1, le=30), user: dict = Depends(current_user)):
    return {"skill": skill, "items": learning_service.store.courses_for_skill(skill, limit)}


@app.get("/api/learning/routes/latest/{analysis_id}")
def latest_learning_route(analysis_id: str, user: dict = Depends(current_user)):
    try:
        route = store.latest_learning_route(analysis_id, user)
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    if not route:
        raise HTTPException(status_code=404, detail="learning route not found")
    return _hydrate_learning_course_covers(route)


@app.get("/api/jobs/facets")
def job_facets():
    jobs = store.list_jobs(active_only=False)
    sources = sorted({job.source_platform for job in jobs if job.source_platform})
    allowed_cities = configured_job_cities()
    cities = [city for city in allowed_cities if any(job.city == city for job in jobs)]
    return {"sources": sources, "cities": cities}


@app.get("/api/knowledge-graph")
def knowledge_graph(
    domain: str | None = None,
    source_platform: str | None = None,
    max_jobs_per_domain: int = Query(default=18, ge=4, le=40),
    max_skills_per_job: int = Query(default=10, ge=3, le=20),
    max_evidence: int = Query(default=360, ge=0, le=1000),
):
    ensure_domain(domain)
    ensure_recruitment_source(source_platform)
    return orchestrator.knowledge_graph_service.build(
        store.jobs.values(), domain=domain, source_platform=source_platform,
        max_jobs_per_domain=max_jobs_per_domain, max_skills_per_job=max_skills_per_job,
        max_evidence=max_evidence,
    )


@app.get("/api/knowledge-graph/career-galaxy")
def career_galaxy(
    domain: str | None = None,
    source_platform: str | None = None,
    max_nodes: int = Query(default=1400, ge=100, le=1800),
):
    ensure_domain(domain)
    ensure_recruitment_source(source_platform)
    return orchestrator.knowledge_graph_service.build_career_galaxy(
        store.jobs.values(),
        domain=domain,
        source_platform=source_platform,
        max_nodes=max_nodes,
        community_signals=store.community_signals.values(),
    )


@app.get("/api/knowledge-graph/relations")
def knowledge_graph_relations(
    domain: str | None = None,
    source_platform: str | None = None,
    period: str | None = None,
):
    ensure_domain(domain)
    ensure_recruitment_source(source_platform)
    return orchestrator.knowledge_graph_service.build_relation_network(
        store.jobs.values(), domain=domain, source_platform=source_platform, period=period,
        emerging_jobs=store.emerging_jobs.values(),
    )


@app.get("/api/knowledge-graph/relations/preview")
def knowledge_graph_relations_preview():
    """Return a small, evidence-backed slice for the public product preview."""
    graph = orchestrator.knowledge_graph_service.build_relation_network(
        store.jobs.values(), emerging_jobs=store.emerging_jobs.values(),
    )
    families = [node for node in graph["nodes"] if node.get("type") == "family"]
    families.sort(key=lambda node: int(node.get("jobCount") or 0), reverse=True)

    selected_families = []
    seen_domains = set()
    for family in families:
        if family.get("domain") not in seen_domains:
            selected_families.append(family)
            seen_domains.add(family.get("domain"))
        if len(selected_families) >= 4:
            break
    for family in families:
        if len(selected_families) >= 4:
            break
        if family not in selected_families:
            selected_families.append(family)

    family_ids = {node["id"] for node in selected_families}
    skill_edges = [
        edge for edge in graph["edges"]
        if edge.get("relationType") in {"core_skill", "requires_skill"}
        and edge.get("source") in family_ids
    ]
    skill_edges.sort(key=lambda edge: (edge.get("source"), -float(edge.get("weight") or 0)))
    selected_skill_edges = []
    family_edge_counts = {}
    for edge in skill_edges:
        family_id = edge["source"]
        if family_edge_counts.get(family_id, 0) >= 3:
            continue
        selected_skill_edges.append(edge)
        family_edge_counts[family_id] = family_edge_counts.get(family_id, 0) + 1
    selected_skill_ids = {edge["target"] for edge in selected_skill_edges}
    selected_skill_pairs = {
        (edge["source"], edge["target"], edge.get("relationType"))
        for edge in selected_skill_edges
    }
    selected_ids = family_ids | selected_skill_ids
    allowed_types = {
        "core_skill", "requires_skill", "job_similarity",
        "skill_cooccurrence", "talent_mobility_inference",
    }
    edges = [
        edge for edge in graph["edges"]
        if edge.get("relationType") in allowed_types
        and edge.get("source") in selected_ids
        and edge.get("target") in selected_ids
        and (
            edge.get("relationType") not in {"core_skill", "requires_skill"}
            or (
                edge.get("source"), edge.get("target"), edge.get("relationType")
            ) in selected_skill_pairs
        )
    ]
    nodes = [node for node in graph["nodes"] if node.get("id") in selected_ids]
    nodes = [{key: value for key, value in node.items() if key != "evidenceIds"} for node in nodes]
    edges = [{key: value for key, value in edge.items() if key != "evidenceIds"} for edge in edges]
    stats = {
        "source_jobs": graph.get("stats", {}).get("source_jobs", 0),
        "nodes": len(nodes), "edges": len(edges),
        "families": len(selected_families), "skills": len(selected_skill_ids),
    }
    return {
        "generated_at": graph.get("generated_at"),
        "data_policy": graph.get("data_policy"),
        "dependency_policy": graph.get("dependency_policy"),
        "period": graph.get("period"), "nodes": nodes, "edges": edges, "stats": stats,
    }


@app.get("/api/graph/overview")
def graph_overview(
    domain: str | None = None,
    source_platform: str | None = None,
):
    ensure_domain(domain)
    ensure_recruitment_source(source_platform)
    return orchestrator.knowledge_graph_service.build_hierarchy_level(
        store.jobs.values(),
        domain=domain,
        source_platform=source_platform,
    )


@app.get("/api/graph/children")
def graph_children(
    parent_id: str,
    domain: str | None = None,
    source_platform: str | None = None,
):
    ensure_domain(domain)
    ensure_recruitment_source(source_platform)
    result = orchestrator.knowledge_graph_service.build_hierarchy_level(
        store.jobs.values(),
        parent_id=parent_id,
        domain=domain,
        source_platform=source_platform,
    )
    if not result:
        raise HTTPException(status_code=404, detail="graph node not found")
    return result


@app.get("/api/graph/temporal")
def graph_temporal(
    domain: str | None = None,
    source_platform: str | None = None,
    max_nodes: int = Query(default=1400, ge=100, le=1800),
):
    ensure_domain(domain)
    ensure_recruitment_source(source_platform)
    return orchestrator.knowledge_graph_service.build_career_galaxy(
        store.jobs.values(),
        domain=domain,
        source_platform=source_platform,
        max_nodes=max_nodes,
        community_signals=store.community_signals.values(),
    )


@app.get("/api/graph/wormhole")
def graph_wormhole(
    domain: str | None = None,
    source_platform: str | None = None,
    max_nodes: int = Query(default=10000, ge=100, le=12000),
):
    """Dense overview with direct job-skill and bounded job-similarity edges."""
    ensure_domain(domain)
    ensure_recruitment_source(source_platform)
    return orchestrator.knowledge_graph_service.build_career_galaxy(
        store.jobs.values(),
        domain=domain,
        source_platform=source_platform,
        max_nodes=max_nodes,
        community_signals=store.community_signals.values(),
        deep_relations=True,
    )


@app.get("/api/graph/nodes/{node_id:path}/evidence")
def graph_node_evidence(node_id: str, limit: int = Query(default=50, ge=1, le=200)):
    relation_graph = orchestrator.knowledge_graph_service.build_relation_network(
        store.jobs.values(),
        emerging_jobs=store.emerging_jobs.values(),
    )
    relation_node = next((item for item in relation_graph["nodes"] if item["id"] == node_id), None)
    if relation_node is not None:
        ids = relation_node.get("evidenceIds", [])
        items = [store.jobs[item_id] for item_id in ids if item_id in store.jobs][:limit]
        return {
            "node": relation_node,
            "total": len(ids),
            "items": items,
            "evidence_policy": "verified_recruitment",
        }
    graph = orchestrator.knowledge_graph_service.build(
        store.jobs.values(),
        max_jobs_per_domain=40,
        max_skills_per_job=20,
        max_evidence=1000,
    )
    node = next((item for item in graph["nodes"] if item["id"] == node_id), None)
    if node is None and node_id in store.jobs:
        job = store.jobs[node_id]
        return {
            "node": {"id": job.source_id, "label": job.title, "type": "job_evidence"},
            "total": 1,
            "items": [job],
            "evidence_policy": "verified_recruitment",
        }
    if node is None:
        raise HTTPException(status_code=404, detail="graph node not found")
    ids = node.get("evidenceIds", [])
    items = [store.jobs[item_id] for item_id in ids if item_id in store.jobs][:limit]
    return {
        "node": node,
        "total": len(ids),
        "items": items,
        "evidence_policy": "verified_recruitment",
    }


@app.get("/api/knowledge-graph/city-map")
def city_talent_map(
    domain: str | None = None,
    source_platform: str | None = None,
    max_cities: int = Query(default=60, ge=8, le=80),
):
    ensure_domain(domain)
    ensure_recruitment_source(source_platform)
    return orchestrator.knowledge_graph_service.build_city_map(
        store.jobs.values(), domain=domain, source_platform=source_platform, max_cities=max_cities,
    )


@app.get("/api/knowledge-graph/city-drilldown")
def city_talent_drilldown(
    city: str = Query(min_length=1, max_length=40),
    domain: str | None = None,
    family: str | None = Query(default=None, max_length=120),
    skill: str | None = Query(default=None, max_length=120),
    source_platform: str | None = None,
):
    ensure_domain(domain)
    ensure_recruitment_source(source_platform)
    payload = orchestrator.knowledge_graph_service.build_city_drilldown(
        store.jobs.values(), city=city, domain=domain, family=family, skill=skill,
        source_platform=source_platform,
    )
    if not payload:
        raise HTTPException(status_code=404, detail="city graph level not found")
    return payload


@app.get("/api/jobs/changes/today")
def today_job_changes():
    changes = store.today_changes()
    return {
        "total": len(changes),
        "summary": {kind: sum(job.change_type == kind for job in changes) for kind in ("new", "updated", "offline")},
        "items": changes,
    }


@app.get("/api/evolution/market")
def market_evolution(
    source_platform: str = Query(default="zhaopin"),
    max_frames: int = Query(default=20, ge=1, le=36),
):
    ensure_recruitment_source(source_platform)
    versions = store.list_job_versions_by_source(source_platform)
    return orchestrator.evolution_agent.analyze_market(
        store.jobs.values(),
        versions,
        source_platform=source_platform,
        max_frames=max_frames,
    )


@app.post("/api/jobs", status_code=201)
def create_job(payload: JobCreateRequest):
    ensure_domain(payload.domain)
    if payload.job_status not in ("active", "offline"):
        raise HTTPException(status_code=422, detail="job_status must be active or offline")
    now = datetime.now(UTC)
    item = CleanedDataItem(
        source_id=payload.source_id or "",
        domain=payload.domain,
        title=payload.title,
        content=payload.content or "；".join(payload.responsibilities),
        collected_at=now,
        quality_score=100,
        dedup_hash=f"enterprise:{payload.company}:{payload.title}",
        company=payload.company,
        city=payload.city,
        district=payload.district,
        salary_min_k=payload.salary_min_k,
        salary_max_k=payload.salary_max_k,
        experience_min_years=payload.experience_min_years,
        experience_max_years=payload.experience_max_years,
        education=payload.education,
        skill_tags=payload.skill_tags,
        responsibilities=payload.responsibilities,
        published_at=now,
        changed_at=now,
        source_platform=payload.source_platform,
        source_url=payload.source_url,
        job_status=payload.job_status,
        change_type="new",
    )
    try:
        return store.create_job(item)
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


@app.get("/api/jobs/{job_id}")
def get_job(job_id: str):
    job = store.get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="job not found")
    return job


@app.get("/api/jobs/{job_id}/versions")
def get_job_versions(job_id: str):
    if store.get_job(job_id) is None:
        raise HTTPException(status_code=404, detail="job not found")
    items = store.list_job_versions(job_id)
    return {"total": len(items), "items": items}


@app.get("/api/jobs/{job_id}/versions/{version_number}/compare")
def compare_job_version(
    job_id: str,
    version_number: int,
    base_version_number: int | None = Query(default=None),
):
    comparison = store.compare_job_version(job_id, version_number, base_version_number)
    if comparison is None:
        raise HTTPException(status_code=404, detail="job version not found")
    return comparison


@app.get("/api/jobs/{job_id}/observations")
def get_job_observations(job_id: str, limit: int = Query(default=100, ge=1, le=500)):
    if store.get_job(job_id) is None:
        raise HTTPException(status_code=404, detail="job not found")
    items = store.list_job_observations(job_id, limit)
    return {"total": len(items), "items": items}


@app.patch("/api/jobs/{job_id}")
def update_job(job_id: str, payload: JobUpdateRequest):
    ensure_domain(payload.domain)
    if payload.job_status not in (None, "active", "offline"):
        raise HTTPException(status_code=422, detail="job_status must be active or offline")
    job = store.update_job(job_id, payload.model_dump(exclude_unset=True))
    if job is None:
        raise HTTPException(status_code=404, detail="job not found")
    return job


@app.delete("/api/jobs/{job_id}", status_code=204)
def delete_job(job_id: str):
    if not store.delete_job(job_id):
        raise HTTPException(status_code=404, detail="job not found")


@app.post("/api/jobs/build-graph")
def build_graph(payload: JobGraphBuildRequest):
    try:
        return orchestrator.build_graph(payload)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/jobs/{job_id}/graph")
def get_graph(job_id: str):
    graph = store.graphs.get(job_id)
    if graph is None:
        raise HTTPException(status_code=404, detail="job graph not found")
    return graph


@app.get("/api/jobs/{job_id}/timeline")
def get_timeline(job_id: str):
    evo = store.evolutions.get(job_id)
    if evo is None:
        try:
            evo = orchestrator.run_evolution(job_id)
        except KeyError as exc:
            raise HTTPException(status_code=404, detail=str(exc)) from exc
    return evo


@app.post("/api/evolution/run")
def run_evolution(payload: EvolutionRunRequest):
    try:
        return orchestrator.run_evolution(payload.job_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.post("/api/assessment/start")
def start_assessment(payload: AssessmentStartRequest, user: dict = Depends(current_user)):
    try:
        session = orchestrator.start_assessment(
            payload.candidate_id,
            payload.target_job_id,
            payload.resume_text,
            payload.resume_parse,
            require_ai_questions=False,
        )
        store.grant_session_access(session, user)
        store.save_candidate_event(session.candidate_id, user, "assessment_started", {"target_job_id": session.target_job_id}, session.session_id)
        store.save_assessment_record(
            {
                "id": session.candidate_id,
                "candidate_id": session.candidate_id,
                "name": "候选人",
                "target_job_id": session.target_job_id,
                "status": "评估进行中",
                "profile": {},
            },
            user,
        )
        return session
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except AssessmentQuestionError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.post("/api/assessment/parse-resume")
def parse_resume(payload: ResumeParseRequest, user: dict = Depends(current_user)):
    if not payload.resume_text.strip():
        raise HTTPException(status_code=422, detail="resume_text is required")
    try:
        parsed = _parse_resume_with_local_fallback(
            payload.resume_text,
            candidate_id=payload.candidate_id,
            source_name=payload.source_name,
        )
        candidate_id = str(parsed.get("candidate_id") or payload.candidate_id or "candidate_unknown")
        store.save_resume_parse_job(candidate_id, {"status": "completed", "resume_parse": parsed, "error": None}, user)
        profile = _profile_from_parse(candidate_id, parsed, payload.resume_text)
        store.save_candidate_profile(profile, user)
        return {**parsed, "candidate_profile": profile.model_dump(mode="json")}
    except ResumeParseError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


def _semantic_enrich_resume(candidate_id: str, resume_text: str, source_name: str, user: dict, text_extraction_ms: int | float = 0):
    started = perf_counter()
    try:
        parsed = _parse_resume_with_local_fallback(
            resume_text,
            candidate_id=candidate_id,
            source_name=source_name,
        )
        provider_status = parsed.get("provider_status", {}) if isinstance(parsed.get("provider_status"), dict) else {}
        parsed_timing = parsed.get("parse_timing", {}) if isinstance(parsed.get("parse_timing"), dict) else {}
        profile = _profile_from_parse(candidate_id, parsed, resume_text)
        resume_parse_jobs[candidate_id] = {
            "status": "completed",
            "resume_parse": parsed,
            "candidate_profile": profile.model_dump(mode="json"),
            "timing": {
                "text_extraction_ms": text_extraction_ms,
                "resume_llm_ms": parsed_timing.get("resume_llm_ms", provider_status.get("duration_ms", 0)),
                "total_ms": round((perf_counter() - started) * 1000) + int(text_extraction_ms or 0),
                "cache_hit": bool(provider_status.get("cache_hit") or parsed_timing.get("cache_hit")),
            },
            "error": None,
        }
        store.save_resume_parse_job(candidate_id, resume_parse_jobs[candidate_id], user)
        store.save_candidate_profile(profile, user)
    except Exception as exc:  # Background enrichment must not invalidate the structured parse.
        current = resume_parse_jobs.get(candidate_id, {})
        resume_parse_jobs[candidate_id] = {
            **current,
            "status": "failed",
            "error": str(exc),
            "timing": {
                **(current.get("timing") or {}),
                "total_ms": round((perf_counter() - started) * 1000) + int(text_extraction_ms or 0),
            },
        }
        store.save_resume_parse_job(candidate_id, resume_parse_jobs[candidate_id], user)


@app.post("/api/assessment/upload-resume")
async def upload_resume(background_tasks: BackgroundTasks, file: UploadFile = File(...), candidate_id: str | None = None, user: dict = Depends(current_user)):
    upload_started = perf_counter()
    resolved_candidate_id = candidate_id or f"candidate_{int(upload_started * 1000)}"
    try:
        extraction_started = perf_counter()
        extracted = extract_resume_text(file.filename or "resume", await file.read())
        text_extraction_ms = round((perf_counter() - extraction_started) * 1000)
        extracted_text = str(extracted["text"])
        images = extracted.get("images") if isinstance(extracted.get("images"), list) and not extracted_text.strip() else None
        parsed = orchestrator.assessment_agent.resume_parser.parse(
            extracted_text,
            candidate_id=resolved_candidate_id,
            source_name=str(extracted["filename"]),
            images=images,
            require_remote=False,
            use_remote=False,
        )
    except ResumeFileError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    except ResumeParseError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    resume_text = extracted_text
    if not resume_text.strip():
        resume_text = str(parsed.get("vision_parse", {}).get("extracted_text", ""))
    provider_status = parsed.get("provider_status", {}) if isinstance(parsed.get("provider_status"), dict) else {}
    parsed_timing = parsed.get("parse_timing", {}) if isinstance(parsed.get("parse_timing"), dict) else {}
    resume_llm_ms = parsed_timing.get("resume_llm_ms", provider_status.get("duration_ms"))
    document_parse = extracted.get("document_parse", {}) if isinstance(extracted.get("document_parse"), dict) else {}
    if parsed.get("vision_parse", {}).get("used"):
        document_parse = {**document_parse, "used_vision": True}
    parsed["document_parse"] = document_parse
    parse_timing = {
        "text_extraction_ms": text_extraction_ms,
        "resume_llm_ms": resume_llm_ms or 0,
        "total_ms": round((perf_counter() - upload_started) * 1000),
        "cache_hit": bool(provider_status.get("cache_hit") or parsed_timing.get("cache_hit")),
    }
    should_enrich = bool(resume_text.strip() and orchestrator.assessment_agent.resume_parser.resume_llm.configured)
    enrichment_status = "processing" if should_enrich else "completed"
    resume_parse_jobs[resolved_candidate_id] = {
        "status": enrichment_status,
        "resume_parse": parsed,
        "timing": parse_timing,
        "error": None,
    }
    store.save_resume_parse_job(resolved_candidate_id, resume_parse_jobs[resolved_candidate_id], user)
    profile = _profile_from_parse(resolved_candidate_id, parsed, resume_text)
    store.save_candidate_profile(profile, user)
    if should_enrich:
        background_tasks.add_task(
            _semantic_enrich_resume,
            resolved_candidate_id,
            resume_text,
            str(extracted["filename"]),
            user,
            text_extraction_ms,
        )
    return {
        "candidate_id": resolved_candidate_id,
        "source_name": extracted["filename"],
        "file_type": extracted["file_type"],
        "text_length": extracted["text_length"],
        "image_count": extracted["image_count"],
        "document_parse": document_parse,
        "resume_text": resume_text,
        "resume_parse": parsed,
        "candidate_profile": profile.model_dump(mode="json"),
        "enrichment_status": enrichment_status,
        "parse_timing": parse_timing,
    }


@app.get("/api/assessment/resume-parses/{candidate_id}")
def get_resume_parse_status(candidate_id: str, user: dict = Depends(current_user)):
    try:
        job = store.get_resume_parse_job(candidate_id, user)
    except PermissionError as exc:
        raise HTTPException(status_code=403, detail=str(exc)) from exc
    if not job:
        raise HTTPException(status_code=404, detail="resume parse job not found")
    return job


@app.get("/api/assessment/sessions/{session_id}")
def get_assessment_session(session_id: str, user: dict = Depends(current_user)):
    session = store.sessions.get(session_id)
    if session is None:
        raise HTTPException(status_code=404, detail="assessment session not found")
    if not store.can_access_session(session_id, user):
        raise HTTPException(status_code=403, detail="assessment session access denied")
    return session


@app.post("/api/assessment/answer")
def answer_assessment(payload: AssessmentAnswer, user: dict = Depends(current_user)):
    if not store.can_access_session(payload.session_id, user):
        raise HTTPException(status_code=403, detail="assessment session access denied")
    try:
        session = orchestrator.answer_assessment(payload)
        store.save_candidate_event(session.candidate_id, user, "dialogue_answered", {"stage": session.current_stage}, session.session_id)
        return session
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except AssessmentQuestionError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc


@app.post("/api/assessment/game-result")
def game_result(payload: GameResultInput, user: dict = Depends(current_user)):
    if not store.can_access_session(payload.session_id, user):
        raise HTTPException(status_code=403, detail="assessment session access denied")
    try:
        session = orchestrator.assessment_game_result(payload)
        store.save_candidate_event(session.candidate_id, user, "game_completed", {"stage": session.current_stage}, session.session_id)
        return session
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.post("/api/assessment/finalize")
def finalize_assessment(payload: AssessmentFinalizeRequest, user: dict = Depends(current_user)):
    if not store.can_access_session(payload.session_id, user):
        raise HTTPException(status_code=403, detail="assessment session access denied")
    try:
        result = orchestrator.finalize_assessment(payload.session_id)
        store.save_candidate_profile(result.candidate_profile, user)
        store.save_candidate_event(result.candidate_profile.candidate_id, user, "assessment_finalized", {"target_job_id": result.candidate_profile.target_job_id}, payload.session_id)
        return result
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.post("/api/match/report")
def match_report(
    payload: MatchReportRequest,
    user: dict = Depends(current_user),
):
    if not store.can_access_candidate(payload.candidate_id, user):
        raise HTTPException(status_code=403, detail="candidate profile access denied")
    try:
        report = orchestrator.match_report(payload.candidate_id, payload.target_job_id)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    profile = store.profiles[report.candidate_id]
    return store.save_report(
        report,
        profile,
        owner_user_id=user["user_id"],
        organization_id=user.get("organization_id"),
    )


def _report_download_response(report, profile, output_format: str) -> StreamingResponse:
    content, media_type, suffix = export_match_report(report, profile, output_format)
    candidate_id = str(report.candidate_id)
    filename = f"讯飞智聘_匹配报告_{candidate_id}.{suffix}"
    ascii_id = re.sub(r"[^A-Za-z0-9._-]+", "_", candidate_id).strip("._") or "candidate"
    disposition = (
        f'attachment; filename="XunfeiZhipin_match_report_{ascii_id}.{suffix}"; '
        f"filename*=UTF-8''{quote(filename, safe='')}"
    )
    headers = {
        "Content-Disposition": disposition,
        "Content-Length": str(len(content)),
    }
    return StreamingResponse(BytesIO(content), media_type=media_type, headers=headers)


def _authorized_report(report_id: str, user: dict) -> MatchReport:
    report = store.reports.get(report_id)
    if report is None:
        raise HTTPException(status_code=404, detail="report not found")
    access = store.report_access.get(report_id)
    if not access:
        raise HTTPException(status_code=403, detail="report access metadata is unavailable")
    is_owner = access.get("owner_user_id") == user["user_id"]
    is_same_organization = (
        user["role"] == "enterprise"
        and bool(access.get("organization_id"))
        and access.get("organization_id") == user.get("organization_id")
    )
    if not (is_owner or is_same_organization):
        raise HTTPException(status_code=403, detail="report access denied")
    return report


@app.post("/api/reports/export")
def export_report(
    payload: ReportExportRequest,
    format: str = Query(default="pdf", pattern="^(pdf|docx)$"),
    user: dict = Depends(current_user),
):
    report = _authorized_report(payload.report_id, user)
    profile = store.report_profiles.get(report.report_id)
    if profile is None:
        raise HTTPException(status_code=404, detail="report profile snapshot not found")
    return _report_download_response(report, profile, format)


@app.get("/api/reports/{report_id}")
def get_report(report_id: str, user: dict = Depends(current_user)):
    return _authorized_report(report_id, user)


@app.get("/api/reports/{report_id}/download")
def download_report(
    report_id: str,
    format: str = Query(default="pdf", pattern="^(pdf|docx)$"),
    user: dict = Depends(current_user),
):
    report = _authorized_report(report_id, user)
    profile = store.report_profiles.get(report.report_id)
    if profile is None:
        raise HTTPException(status_code=404, detail="report profile snapshot not found")
    return _report_download_response(report, profile, format)


@app.get("/api/candidates/{candidate_id}/profile")
def get_candidate_profile(candidate_id: str, authorization: str | None = Header(default=None)):
    profile = store.profiles.get(candidate_id)
    if profile is None:
        raise HTTPException(status_code=404, detail="candidate profile not found")
    user = auth_service.resolve_token(_bearer_token(authorization))
    if not user:
        raise HTTPException(status_code=401, detail="请先登录后继续")
    if not store.can_access_candidate(candidate_id, user):
        raise HTTPException(status_code=403, detail="candidate profile access denied")
    return profile


@app.post("/api/emerging-jobs/discover")
def discover_emerging_jobs(payload: EmergingJobDiscoveryRequest):
    try:
        items = orchestrator.discover_emerging_jobs(payload)
    except KeyError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    return {"total": len(items), "items": items}


@app.post("/api/emerging-jobs/runs")
def run_emerging_job_discovery(payload: EmergingJobDiscoveryRequest):
    started = perf_counter()
    items = orchestrator.discover_emerging_jobs(payload)
    return {
        "run_id": f"emerging_run_{datetime.now(UTC):%Y%m%d%H%M%S}",
        "status": "completed",
        "duration_ms": round((perf_counter() - started) * 1000),
        "total": len(items),
        "items": items,
        "evidence_policy": {
            "boss_zhaopin": "formal_job_evidence",
            "community_sources": "auxiliary_trend_only",
        },
    }


@app.get("/api/emerging-jobs")
def list_emerging_jobs(
    domain: str | None = None,
    maturity_status: str | None = None,
    minimum_confidence: float = Query(default=0, ge=0, le=1),
):
    items = list(store.emerging_jobs.values())
    if domain:
        items = [item for item in items if item.domain == domain]
    if maturity_status:
        if maturity_status not in {"signal", "forming", "established"}:
            raise HTTPException(status_code=422, detail="invalid maturity_status")
        items = [item for item in items if item.maturity_status == maturity_status]
    items = [item for item in items if item.confidence >= minimum_confidence]
    items.sort(key=lambda item: (item.confidence, item.growth_rate), reverse=True)
    return {"total": len(items), "items": items}


@app.post("/api/emerging-jobs/{candidate_id}/review")
def review_emerging_job(candidate_id: str, payload: EmergingJobReviewRequest):
    if payload.status not in {"pending", "watching", "approved", "rejected"}:
        raise HTTPException(status_code=422, detail="invalid emerging job review status")
    candidate = store.review_emerging_job(candidate_id, payload.status, payload.corrections)
    if candidate is None:
        raise HTTPException(status_code=404, detail="emerging job candidate not found")
    return {"candidate": candidate, "note": payload.note}


@app.post("/api/validation/runs", status_code=201)
def create_validation_run(payload: ValidationRunRequest):
    return orchestrator.run_validation(payload)


@app.get("/api/validation/runs")
def list_validation_runs(limit: int = Query(default=20, ge=1, le=100)):
    runs = sorted(store.validation_runs.values(), key=lambda run: run.created_at, reverse=True)[:limit]
    return {
        "total": len(store.validation_runs),
        "items": [
            {
                "run_id": run.run_id,
                "status": run.status,
                "stage": run.stage,
                "stage_progress": run.stage_progress,
                "metrics": run.metrics,
                "failure_count": len(run.failures),
                "duration_ms": run.duration_ms,
                "version": run.version,
                "created_at": run.created_at,
                "completed_at": run.completed_at,
            }
            for run in runs
        ],
    }


@app.get("/api/validation/runs/{run_id}")
def get_validation_run(run_id: str):
    run = store.validation_runs.get(run_id)
    if run is None:
        raise HTTPException(status_code=404, detail="validation run not found")
    return run


@app.get("/api/demo/state")
def demo_state():
    return {
        "batches": list(store.batches),
        "graphs": list(store.graphs),
        "evolutions": list(store.evolutions),
        "sessions": list(store.sessions),
        "profiles": list(store.profiles),
        "reports": list(store.reports),
        "jobs": len(store.jobs),
        "today_changes": len(store.today_changes()),
    }
