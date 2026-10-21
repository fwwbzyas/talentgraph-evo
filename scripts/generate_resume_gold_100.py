"""Generate a deterministic, privacy-free Chinese resume benchmark."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as pdf_canvas
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "data/evaluation/resume_gold_100"
SEED = 20260824
DOMAINS = ("人工智能", "大数据", "智能系统", "物联网")
SPLITS = ("calibration", "regression", "blind")
LEVEL_COUNTS = {"初级": 25, "中级": 30, "高级": 45}
DIFFICULTY_COUNTS = {"易": 20, "中": 40, "难": 40}

NAMES = (
    "林知微", "周予安", "沈闻舟", "顾清和", "许星河", "唐映雪", "程远川",
    "苏晚晴", "陆景深", "叶初阳", "韩致远", "江若宁", "白承宇", "邵明哲",
    "方若琳", "严书航", "贺云开", "罗以宁", "温景行", "秦思源", "齐墨",
    "梁书瑶", "傅斯年", "郑南乔", "钟亦航", "林亦辰", "周沐言", "沈嘉树",
    "顾言川", "许清越", "唐子衿", "程砚秋", "苏念安", "陆知行", "叶舒然",
    "韩星野", "江临川", "白芷宁", "邵言蹊", "方时雨", "严知远", "贺清屿",
    "罗嘉禾", "温以辰", "秦若川", "齐书宁", "梁景明", "傅予墨", "郑语嫣",
    "钟云舒", "林溪", "周砚", "沈青禾", "顾南风", "许澄", "唐昀",
    "程知夏", "苏言默", "陆昭", "叶青宁", "韩屿", "江临安", "白闻远",
    "邵锦程", "方知意", "严澜", "贺知秋", "罗星辰", "温书言", "秦望舒",
    "齐修远", "梁清妍", "傅景舟", "郑知遥", "钟予希", "林若川", "周清野",
    "沈星言", "顾书宁", "许临渊", "唐安", "程砚", "苏云深", "陆清嘉",
    "叶知秋", "韩予辰", "江晚舟", "白清越", "邵景行", "方予安", "严知微",
    "贺言川", "罗清和", "温予墨", "秦星河", "齐映雪", "梁远川", "傅晚晴",
    "郑景深", "钟初阳", "林致远", "周若宁",
)

LOCATIONS = ("北京", "上海", "深圳", "杭州", "广州", "成都", "南京", "武汉", "西安", "远程")
SCHOOLS = ("东澜理工大学", "南岭科技大学", "海川大学", "北辰信息学院", "云岫工业大学")
COMPANY_PREFIXES = ("澄原", "远川", "星桥", "云栖", "青岚", "启明", "瀚海", "知行", "栖木", "曜石")
SCENES = (
    "客服知识分散在工单、手册和历史对话中，业务希望缩短一线检索时间",
    "设备运行数据增长快，原有批处理无法及时发现异常",
    "多团队使用同一数据口径，指标定义不一致影响经营判断",
    "生产现场存在低频故障，人工巡检难以持续覆盖关键环节",
    "模型上线后反馈闭环较慢，样本质量和服务稳定性需要同步治理",
    "跨区域终端网络质量波动，边缘节点需要在断网时维持基本能力",
    "实验室方案即将进入真实业务，必须把原型转成可观测、可回滚的服务",
    "管理者希望看到可解释的预测依据，而不是只有一个分数",
    "历史系统接口多且标准不一，新增能力不能破坏已有流程",
    "高峰期任务集中到达，资源调度和服务响应需要可量化优化",
)
WORK_STYLES = (
    "先拆解口径再动手实现，习惯用小样本和监控指标验证判断",
    "擅长把业务描述翻译成接口、数据契约和可复用组件",
    "面对不确定需求时会先明确边界，再通过灰度迭代收敛方案",
    "重视上线后的反馈闭环，会把故障复盘转成自动化检查项",
    "能够在研发速度、可维护性和交付风险之间做取舍并留下记录",
)

ROLE_SPECS: tuple[dict[str, Any], ...] = (
    {"domain": "人工智能", "role": "AI应用工程师", "family": "AI应用与检索增强", "focus": "知识问答、智能助手和模型服务落地", "major": "人工智能", "groups": {"编程语言": ["Python", "SQL", "Git"], "模型与算法": ["LLM", "RAG", "Prompt工程"], "服务与工程": ["FastAPI", "Redis", "Docker"], "质量与观测": ["Pytest", "Prometheus", "日志追踪"]}, "actions": ["用分层召回、重排和引用约束降低回答漂移", "把异步任务、缓存和限流封装为可观测接口", "用离线评测集和线上反馈分开验证召回与生成效果"], "nouns": ["知识问答工作台", "售后助手", "研发资料检索服务", "智能工单归因"], "metrics": ["P95延迟从 {a}ms 降至 {b}ms", "引用命中率提升 {p}%", "人工复核量下降 {q}%"]},
    {"domain": "人工智能", "role": "机器学习工程师", "family": "预测建模与特征工程", "focus": "预测模型、特征服务和实验迭代", "major": "计算机科学与技术", "groups": {"编程语言": ["Python", "SQL", "Git"], "模型与算法": ["LightGBM", "PyTorch", "特征工程"], "训练与服务": ["MLflow", "FastAPI", "Docker"], "质量与观测": ["交叉验证", "Prometheus", "模型监控"]}, "actions": ["建立特征版本和训练数据校验，避免离线线上口径漂移", "用分层抽样、校准曲线和阈值实验确定发布方案", "将推理服务拆出批量与实时路径并补齐回滚开关"], "nouns": ["需求预测模型", "风险评分服务", "设备寿命预测", "客户流失预警"], "metrics": ["AUC提升 {p}%", "召回率提升 {q}%", "误报率下降 {r}%"]},
    {"domain": "人工智能", "role": "NLP算法工程师", "family": "自然语言处理", "focus": "文本理解、信息抽取和语义检索", "major": "软件工程", "groups": {"编程语言": ["Python", "SQL", "Git"], "模型与算法": ["Transformer", "BERT", "文本分类"], "数据与服务": ["标注规范", "Elasticsearch", "FastAPI"], "质量与观测": ["F1评估", "错误分析", "Docker"]}, "actions": ["按业务标签边界整理困难样本并设计可回溯标注规则", "组合词法特征与预训练模型，针对长文本做分段聚合", "用混淆矩阵和人工抽检定位召回损失，再迭代阈值"], "nouns": ["合同条款抽取", "舆情主题识别", "工单语义分流", "政策文本检索"], "metrics": ["Macro-F1提升 {p}%", "平均处理时长下降 {q}%", "长尾类别召回提升 {r}%"]},
    {"domain": "人工智能", "role": "计算机视觉工程师", "family": "视觉识别与质检", "focus": "视觉检测、图像数据闭环和边缘推理", "major": "电子信息工程", "groups": {"编程语言": ["Python", "C++", "Git"], "模型与算法": ["PyTorch", "目标检测", "OpenCV"], "训练与部署": ["ONNX", "TensorRT", "Docker"], "质量与观测": ["数据增强", "mAP评估", "模型监控"]}, "actions": ["按光照、角度和缺陷类型拆分样本并维护数据版本", "通过增强、难例挖掘和阈值分层提高小目标召回", "将模型导出为ONNX并在边缘设备上完成性能剖析"], "nouns": ["产线缺陷检测", "仓储包裹识别", "票据图像审核", "道路目标感知"], "metrics": ["mAP提升 {p}%", "单帧延迟从 {a}ms 降至 {b}ms", "漏检率下降 {q}%"]},
    {"domain": "大数据", "role": "数据工程师", "family": "离线数仓与数据管道", "focus": "数仓建模、数据管道和可复用数据服务", "major": "数据科学与大数据技术", "groups": {"编程与查询": ["Python", "SQL", "Shell"], "计算与存储": ["Spark", "Hive", "对象存储"], "建模与调度": ["维度建模", "Airflow", "数据分层"], "质量与协作": ["Great Expectations", "Git", "Docker"]}, "actions": ["按事实、维度和口径拆解宽表，减少下游重复加工", "为增量任务设计幂等键、分区策略和失败重跑路径", "将数据质量规则接入调度并把异常样本回流给业务确认"], "nouns": ["经营分析数仓", "供应链主题数据集市", "用户行为宽表", "财务对账管道"], "metrics": ["日处理数据量达到 {n}TB", "任务耗时下降 {q}%", "数据新鲜度提升至 {h} 分钟"]},
    {"domain": "大数据", "role": "实时计算工程师", "family": "流式计算与实时数仓", "focus": "事件流处理、实时指标和稳定性治理", "major": "计算机科学与技术", "groups": {"编程与查询": ["Java", "SQL", "Python"], "流式计算": ["Flink", "Kafka", "实时数仓"], "工程与存储": ["Checkpoint", "Iceberg", "Docker"], "质量与观测": ["Prometheus", "Grafana", "压测"]}, "actions": ["依据事件时间和水位线处理迟到数据，控制状态膨胀", "用分区、并行度和反压指标定位高峰期吞吐瓶颈", "建立重放、旁路和告警策略，让异常恢复可演练"], "nouns": ["实时风控流水线", "订单状态看板", "物流轨迹计算", "营销事件平台"], "metrics": ["吞吐达到 {n}万条/秒", "端到端延迟下降 {q}%", "故障恢复时间缩短 {r}%"]},
    {"domain": "大数据", "role": "数据分析师", "family": "经营分析与实验评估", "focus": "指标体系、专题分析和实验决策支持", "major": "统计学", "groups": {"分析语言": ["SQL", "Python", "R"], "分析方法": ["指标体系", "漏斗分析", "A/B测试"], "数据工具": ["Spark", "Tableau", "数据可视化"], "质量与协作": ["统计检验", "数据核验", "业务沟通"]}, "actions": ["先锁定指标定义、统计窗口和样本范围，再形成分析结论", "用分群、漏斗和留存拆解变化来源，区分相关与因果", "将结论沉淀为看板和行动清单，跟踪上线后的真实反馈"], "nouns": ["用户增长分析", "供应链效率诊断", "内容推荐复盘", "门店经营看板"], "metrics": ["分析周期缩短 {q}%", "关键指标提升 {p}%", "异常发现提前 {d} 天"]},
    {"domain": "大数据", "role": "数据治理工程师", "family": "数据质量与资产治理", "focus": "数据标准、质量规则和资产可发现性", "major": "信息管理与信息系统", "groups": {"编程与查询": ["Python", "SQL", "Shell"], "治理方法": ["数据标准", "元数据管理", "血缘分析"], "平台与质量": ["Great Expectations", "Airflow", "数据目录"], "工程与协作": ["API设计", "Git", "故障复盘"]}, "actions": ["把关键字段、责任人和质量阈值登记为可执行规则", "沿数据血缘追踪异常影响面，优先治理高频断点", "设计质量评分、申诉和复核流程，减少口径争议"], "nouns": ["主数据治理中心", "数据质量门户", "指标资产目录", "供应商数据校验"], "metrics": ["核心表覆盖率达到 {p}%", "质量告警误报下降 {q}%", "问题平均闭环时间缩短 {r}%"]},
    {"domain": "智能系统", "role": "智能系统工程师", "family": "智能调度与系统集成", "focus": "智能调度、服务编排和复杂系统集成", "major": "自动化", "groups": {"编程与算法": ["Python", "C++", "算法设计"], "系统工程": ["Linux", "消息队列", "REST API"], "部署与观测": ["Docker", "Kubernetes", "Prometheus"], "可靠性": ["状态机", "故障演练", "接口测试"]}, "actions": ["把业务流程建模为状态机，明确超时、重试和人工接管边界", "用事件驱动方式拆解模块，降低跨系统调用的耦合", "通过压测、故障注入和链路观测验证高峰期行为"], "nouns": ["园区智能调度", "仓内任务编排", "能源优化控制", "服务机器人中台"], "metrics": ["调度成功率提升 {p}%", "平均响应时延下降 {q}%", "人工干预次数减少 {r}%"]},
    {"domain": "智能系统", "role": "机器人感知工程师", "family": "机器人感知与定位", "focus": "多传感器融合、定位和环境感知", "major": "机器人技术", "groups": {"编程与算法": ["C++", "Python", "ROS2"], "感知算法": ["SLAM", "点云处理", "传感器融合"], "系统与部署": ["Linux", "OpenCV", "Docker"], "质量与观测": ["回放评测", "仿真测试", "故障诊断"]}, "actions": ["对齐时间戳和坐标系，建立传感器标定与回放工具链", "融合激光、视觉和里程计信息，针对遮挡场景调整策略", "在仿真与实机分别评估漂移、丢帧和异常恢复能力"], "nouns": ["室内配送机器人", "仓储叉车定位", "巡检机器人感知", "园区移动平台"], "metrics": ["定位漂移下降 {q}%", "障碍物召回率达到 {p}%", "实机异常恢复时间缩短 {r}%"]},
    {"domain": "智能系统", "role": "MLOps平台工程师", "family": "模型平台与智能运维", "focus": "模型交付、资源调度和机器学习平台稳定性", "major": "软件工程", "groups": {"编程与工程": ["Python", "Go", "Git"], "平台能力": ["Kubernetes", "Docker", "MLflow"], "服务与数据": ["FastAPI", "Kafka", "对象存储"], "质量与观测": ["Prometheus", "灰度发布", "CI/CD"]}, "actions": ["定义训练、评测、发布和回滚的流水线状态与权限边界", "按资源队列和优先级调度任务，降低空闲与抢占浪费", "补齐模型版本、特征版本和服务指标的关联追踪"], "nouns": ["模型交付平台", "特征服务中台", "训练资源调度", "智能运维门户"], "metrics": ["发布周期从 {a} 小时降至 {b} 小时", "资源利用率提升 {p}%", "失败任务重试率下降 {q}%"]},
    {"domain": "智能系统", "role": "智能制造解决方案工程师", "family": "智能制造与现场交付", "focus": "产线系统方案、现场交付和持续改进", "major": "工业工程", "groups": {"编程与数据": ["Python", "SQL", "JavaScript"], "工业系统": ["MES", "OPC UA", "设备数据"], "方案与交付": ["流程建模", "接口集成", "现场调试"], "质量与协作": ["问题闭环", "验收测试", "项目管理"]}, "actions": ["把现场工艺、设备信号和质量要求整理为可验收流程", "设计接口映射和异常补偿，保证旧设备接入不影响生产", "通过分阶段上线、现场培训和复盘降低交付波动"], "nouns": ["产线质量追溯", "设备协同排产", "能耗优化系统", "工厂异常闭环"], "metrics": ["换线时间缩短 {q}%", "关键设备可视率达到 {p}%", "异常闭环周期下降 {r}%"]},
    {"domain": "物联网", "role": "物联网平台工程师", "family": "设备接入与平台服务", "focus": "设备接入、数字孪生和平台服务治理", "major": "物联网工程", "groups": {"编程与工程": ["Java", "Python", "Git"], "连接与协议": ["MQTT", "HTTP", "设备影子"], "平台与存储": ["Redis", "PostgreSQL", "Docker"], "质量与观测": ["Prometheus", "契约测试", "限流熔断"]}, "actions": ["抽象设备模型、属性、事件和命令，减少厂商差异向上扩散", "设计断线重连、幂等消费和离线补偿，稳定设备状态同步", "通过租户隔离、指标看板和灰度策略控制平台风险"], "nouns": ["园区设备接入平台", "冷链设备云", "智慧楼宇数字孪生", "资产远程运维"], "metrics": ["日接入设备达到 {n} 万台", "消息处理延迟下降 {q}%", "设备在线率提升 {p}%"]},
    {"domain": "物联网", "role": "嵌入式物联网工程师", "family": "嵌入式设备与固件", "focus": "设备固件、低功耗通信和可靠性验证", "major": "电子信息工程", "groups": {"语言与嵌入式": ["C", "C++", "FreeRTOS"], "连接与协议": ["BLE", "LoRaWAN", "UART"], "硬件与调试": ["STM32", "I2C", "逻辑分析仪"], "质量与交付": ["单元测试", "OTA升级", "故障复现"]}, "actions": ["拆分驱动、业务状态和升级模块，降低硬件变更影响面", "用功耗采样、边界测试和异常注入验证长时间运行稳定性", "设计分片传输、断点续传和失败回滚，保障远程升级"], "nouns": ["低功耗环境采集器", "智能门锁固件", "工业传感终端", "可穿戴设备"], "metrics": ["待机功耗下降 {q}%", "升级成功率达到 {p}%", "现场返修率下降 {r}%"]},
    {"domain": "物联网", "role": "边缘计算工程师", "family": "边缘推理与节点管理", "focus": "边缘节点、离线计算和模型推理优化", "major": "计算机科学与技术", "groups": {"编程与工程": ["Python", "C++", "Linux"], "边缘与模型": ["Docker", "ONNX Runtime", "边缘推理"], "连接与同步": ["MQTT", "gRPC", "断点同步"], "质量与观测": ["资源隔离", "Prometheus", "现场日志"]}, "actions": ["将采集、推理和上报拆成可独立降级的边缘服务", "通过量化、批处理和资源配额控制节点CPU与内存峰值", "设计离线队列、配置签名和远程诊断，减少现场往返"], "nouns": ["边缘质检节点", "园区视频分析", "矿区安全感知", "零售门店边缘盒"], "metrics": ["端侧延迟下降 {q}%", "带宽占用下降 {r}%", "离线缓存恢复率达到 {p}%"]},
    {"domain": "物联网", "role": "IoT数据工程师", "family": "时序数据与设备分析", "focus": "时序数据管道、设备指标和异常分析", "major": "数据科学与大数据技术", "groups": {"编程与查询": ["Python", "SQL", "Shell"], "时序与计算": ["InfluxDB", "Flink", "时序特征"], "连接与平台": ["Kafka", "MQTT", "数据湖"], "质量与观测": ["数据校验", "Grafana", "异常检测"]}, "actions": ["按设备、测点和采样周期设计时序数据模型与保留策略", "处理乱序、补点和重复上报，保证分析窗口口径稳定", "把异常规则与维修工单关联，验证告警是否真正可行动"], "nouns": ["设备健康分析", "冷链温度监控", "能源时序平台", "机房环境预警"], "metrics": ["数据完整率提升 {p}%", "告警提前 {d} 小时", "存储成本下降 {q}%"]},
)

LAYOUTS = (
    {"name": "经典蓝", "hex": "1F4E79", "pdf": colors.HexColor("#1F4E79")},
    {"name": "极简灰", "hex": "4B5563", "pdf": colors.HexColor("#4B5563")},
    {"name": "湖水青", "hex": "087F8C", "pdf": colors.HexColor("#087F8C")},
    {"name": "工程橙", "hex": "B45309", "pdf": colors.HexColor("#B45309")},
    {"name": "森林绿", "hex": "276749", "pdf": colors.HexColor("#276749")},
    {"name": "石墨黑", "hex": "27272A", "pdf": colors.HexColor("#27272A")},
    {"name": "深红色", "hex": "9B2C2C", "pdf": colors.HexColor("#9B2C2C")},
    {"name": "靛青色", "hex": "3730A3", "pdf": colors.HexColor("#3730A3")},
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _flatten_skills(groups: dict[str, list[str]]) -> list[str]:
    return [skill for values in groups.values() for skill in values]


def _level_sequence(seed: int) -> list[str]:
    values = [level for level, count in LEVEL_COUNTS.items() for _ in range(count)]
    for pos in range(len(values) - 1, 0, -1):
        swap = (seed * 17 + pos * 31) % (pos + 1)
        values[pos], values[swap] = values[swap], values[pos]
    return values


def _difficulty_sequence(seed: int) -> list[str]:
    values = [level for level, count in DIFFICULTY_COUNTS.items() for _ in range(count)]
    for pos in range(len(values) - 1, 0, -1):
        swap = (seed * 13 + pos * 19) % (pos + 1)
        values[pos], values[swap] = values[swap], values[pos]
    return values


def _duration(start: str, end: str) -> float:
    sy, sm = (int(value) for value in start.split("."))
    ey, em = (int(value) for value in end.split("."))
    return round(((ey - sy) * 12 + em - sm) / 12, 1)


def _timeline(index: int, level: str, spec: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], float]:
    offset = index % 3
    if level == "初级":
        graduate = 2023 + index % 2
        topic = re.sub(r"[，,]+$", "", SCENES[index % len(SCENES)][:12])
        education = [{"institution": SCHOOLS[index % len(SCHOOLS)], "degree": "本科", "major": spec["major"], "start": f"{graduate - 4}.09", "end": f"{graduate}.06", "coursework": ["数据结构", "数据库原理", "软件工程"], "capstone": f"面向{topic}的应用研究"}]
        work = [{"start": f"{graduate}.07", "end": f"{graduate + 1}.02", "company": f"{COMPANY_PREFIXES[index % len(COMPANY_PREFIXES)]}技术实验室", "title": f"{spec['role']}实习生", "scope": "在导师指导下负责模块实现、数据整理和测试闭环", "responsibilities": ["拆解需求并完成可验收的小功能", "记录实验参数、异常样本和复盘结论"], "achievements": [f"将一个关键流程的人工操作减少约{10 + index % 9}%", "补齐接口测试和使用说明，便于后续交接"]}]
    elif level == "中级":
        graduate = 2018 + offset
        education = [{"institution": SCHOOLS[(index + 1) % len(SCHOOLS)], "degree": "本科", "major": spec["major"], "start": f"{graduate - 4}.09", "end": f"{graduate}.06", "coursework": ["分布式系统", "统计学习", "项目管理"], "capstone": f"面向{spec['focus']}的工程化实践"}]
        work = [
            {"start": f"{graduate}.07", "end": f"{graduate + 2}.01", "company": f"{COMPANY_PREFIXES[index % len(COMPANY_PREFIXES)]}数智科技", "title": spec["role"], "scope": "负责核心模块开发、数据口径确认和上线验收", "responsibilities": ["与产品和业务确认输入、边界及验收指标", "维护关键接口、任务链路和异常处理"], "achievements": [f"支持{3 + index % 5}个业务团队稳定使用", f"通过专项治理将返工率降低{12 + index % 12}%"]},
            {"start": f"{graduate + 2}.02", "end": f"{graduate + 5}.06", "company": f"{COMPANY_PREFIXES[(index + 3) % len(COMPANY_PREFIXES)]}智能服务", "title": spec["role"], "scope": "承担子系统方案、技术评审和跨团队交付推进", "responsibilities": ["拆分迭代任务并为关键路径补充观测指标", "协同测试和运维完成灰度、回滚与复盘"], "achievements": [f"把核心链路稳定性提升到{96 + index % 3}.{index % 10}%", f"沉淀{4 + index % 5}份可复用规范或工具"]},
        ]
    else:
        graduate = 2014 + offset
        education = [{"institution": SCHOOLS[(index + 2) % len(SCHOOLS)], "degree": "本科", "major": spec["major"], "start": f"{graduate - 7}.09", "end": f"{graduate - 3}.06", "coursework": ["高级算法", "系统架构", "质量工程"], "capstone": f"面向{spec['focus']}的系统设计"}, {"institution": SCHOOLS[(index + 3) % len(SCHOOLS)], "degree": "硕士", "major": spec["major"], "start": f"{graduate - 2}.09", "end": f"{graduate}.06", "coursework": ["优化方法", "工程统计", "技术管理"], "capstone": f"{spec['family']}中的可靠性与效率研究"}]
        work = [
            {"start": f"{graduate}.07", "end": f"{graduate + 3}.08", "company": f"{COMPANY_PREFIXES[index % len(COMPANY_PREFIXES)]}信息系统", "title": f"高级{spec['role']}", "scope": "负责基础模块建设、疑难问题定位和交付标准制定", "responsibilities": ["把复杂需求拆成可验证的架构和接口约束", "指导成员完成实现、测试与上线复盘"], "achievements": [f"支撑年增长约{18 + index % 12}%的业务规模", "建立从设计评审到故障复盘的质量机制"]},
            {"start": f"{graduate + 3}.09", "end": f"{graduate + 7}.12", "company": f"{COMPANY_PREFIXES[(index + 4) % len(COMPANY_PREFIXES)]}数据科技", "title": f"技术负责人（{spec['role']}）", "scope": "负责领域方案、关键技术决策和多团队协同", "responsibilities": ["制定版本节奏、风险清单和跨系统依赖方案", "评审关键代码和数据变更，保证迁移可回退"], "achievements": [f"管理{3 + index % 4}条核心链路并保持稳定运行", f"通过架构调整将资源成本降低{15 + index % 16}%"]},
            {"start": f"{graduate + 8}.01", "end": "2026.06", "company": f"{COMPANY_PREFIXES[(index + 6) % len(COMPANY_PREFIXES)]}智能平台", "title": f"资深{spec['role']}", "scope": "负责平台化演进、技术路线和关键项目结果", "responsibilities": ["定义领域能力边界，推动组件、规范和指标统一", "在资源受限场景下做取舍并保障关键目标交付"], "achievements": [f"服务{20 + index % 18}个内部项目或生产场景", f"把关键故障平均恢复时间缩短{24 + index % 18}%"]},
        ]
    experience = round(sum(_duration(item["start"], item["end"]) for item in work), 1)
    return education, work, experience


def _metric_values(index: int) -> dict[str, int]:
    return {"a": 360 + index % 90, "b": 180 + index % 70, "p": 12 + index % 17, "q": 15 + index % 19, "r": 18 + index % 16, "n": 2 + index % 8, "h": 5 + index % 11, "d": 1 + index % 5}


def _projects(index: int, level: str, spec: dict[str, Any], work: list[dict[str, Any]]) -> list[dict[str, Any]]:
    values = _metric_values(index)
    count = 3 if level == "高级" else 2
    projects: list[dict[str, Any]] = []
    for number in range(count):
        scene = SCENES[(index * 2 + number) % len(SCENES)]
        noun = spec["nouns"][(index + number) % len(spec["nouns"])]
        period = work[min(number, len(work) - 1)]
        project_values = {key: value + number * (3 if key in {"p", "q", "r"} else 24) for key, value in values.items()}
        raw_metrics = [metric.format(**project_values) for metric in spec["metrics"]]
        action = [spec["actions"][(index + number + step) % len(spec["actions"])] for step in range(3)]
        background = f"{scene}。项目目标是建设{noun}，并让结果能够被业务团队稳定复用。"
        boundary = f"本人负责{spec['focus']}中的{noun}方案拆解、核心实现和上线验证；与产品共同确认口径，与测试和运维协作交付，不替代硬件采购、视觉设计或业务审批职责。"
        result = f"上线后形成可复盘的指标闭环：{raw_metrics[0]}，同时{raw_metrics[1]}，并通过灰度和回滚演练确认{raw_metrics[2]}。"
        projects.append({
            "name": f"{noun}·第{index + number + 1}期",
            "period": f"{period['start']}-{period['end']}",
            "background": background,
            "responsibility_boundary": boundary,
            "technical_actions": action,
            "quantified_results": raw_metrics,
            "situation": background,
            "task": f"在既有系统和明确资源约束下，交付{noun}的可用版本并达到验收指标。",
            "action": action,
            "result": result,
            "tech_stack": _flatten_skills(spec["groups"])[number:number + 6],
            "metrics": raw_metrics,
            "role": "核心负责人" if level == "高级" else "核心开发",
        })
    return projects


def _skill_evidence(index: int, spec: dict[str, Any], projects: list[dict[str, Any]], work: list[dict[str, Any]]) -> list[dict[str, str]]:
    evidence: list[dict[str, str]] = []
    for group, skills in spec["groups"].items():
        for offset, skill in enumerate(skills):
            project = projects[(index + offset) % len(projects)]
            action = project["technical_actions"][offset % len(project["technical_actions"])]
            if action.startswith("用"):
                action_sentence = f"使用{action[1:]}"
            elif action.startswith(("把", "将")):
                action_sentence = f"设计并实现{action[1:]}"
            else:
                action_sentence = action
            evidence.append({"skill": skill, "group": group, "evidence": f"在《{project['name']}》中，{action_sentence}，最终{project['quantified_results'][offset % len(project['quantified_results'])]}。", "source": "项目经历", "confidence": "高"})
    return evidence


def _profile(index: int, level: str, spec: dict[str, Any], experience: float) -> dict[str, Any]:
    name = NAMES[index % len(NAMES)]
    location = LOCATIONS[(index * 3) % len(LOCATIONS)]
    return {"candidate_id": f"resume_{index + 1:03d}", "name": name, "fictional": True, "privacy_note": "虚构匿名样本，不含电话、邮箱、身份证号或真实联系地址", "preferred_locations": [location, "远程"] if index % 3 == 0 else [location], "availability": "一个月内到岗" if level != "高级" else "双方确认后到岗", "experience_years": experience, "role_family": spec["family"], "domain": spec["domain"], "level": level}


def _summary(index: int, level: str, spec: dict[str, Any], experience: float) -> str:
    style = WORK_STYLES[index % len(WORK_STYLES)]
    prefix = "具备扎实项目训练的应届及早期工程师" if level == "初级" else "能够独立负责模块并推动跨团队交付的工程师" if level == "中级" else "能够做技术取舍、建立工程机制并带动团队交付的资深工程师"
    variants = (
        f"{prefix}，工作主线围绕{spec['focus']}，累计相关经历约{experience:.1f}年。{style}，能够在{spec['family']}场景中完成从问题澄清到上线复盘的闭环。",
        f"{prefix}，专注{spec['focus']}，有约{experience:.1f}年的相关实践。{style}。曾在{spec['family']}项目中推动方案落地，重视结果可解释、系统可维护和指标可复现。",
        f"{prefix}，将{spec['focus']}作为主要能力方向，累计经历约{experience:.1f}年。{style}。熟悉从需求边界、技术实现到效果验证的完整链路，并能在{spec['family']}场景中持续复盘。",
        f"{prefix}，围绕{spec['focus']}积累了约{experience:.1f}年经验。{style}。在{spec['family']}场景中关注数据口径、工程质量和业务反馈，能把阶段性方案沉淀为可复用能力。",
        f"{prefix}，主要投入{spec['focus']}方向，相关经历约{experience:.1f}年。{style}。擅长在{spec['family']}项目里明确协作边界，以可观测指标和复盘记录支撑技术决策。",
    )
    return variants[index % len(variants)]


def _strengths_weaknesses(index: int, level: str, spec: dict[str, Any]) -> tuple[list[str], list[str]]:
    strength_variants = (
        f"问题拆解清晰：能将{spec['focus']}拆成数据、接口、验证和交付四类可跟踪任务。",
        f"落地节奏稳定：面对{spec['focus']}需求，会先做最小可验收版本，再用指标推动迭代。",
        f"工程意识较强：能把{spec['family']}中的隐性风险转成接口约束、检查项和复盘记录。",
        f"沟通表达具体：习惯用输入、边界和结果描述{spec['focus']}方案，便于不同角色快速对齐。",
        f"学习迁移能力较好：能将一次{spec['family']}项目中的经验整理为下一次交付可复用的实践。",
    )
    evidence_variants = (
        f"证据意识较强：习惯用项目指标、错误样本和复盘记录说明{spec['family']}方案是否有效。",
        f"验证习惯扎实：会把{spec['family']}的预期结果拆成离线样例、线上指标和人工抽检三层证据。",
        f"重视可复现性：对{spec['focus']}的关键结论会保留数据口径、实验条件和版本信息。",
        f"结果导向明确：不仅关注功能完成，也会追踪{spec['family']}上线后的使用反馈和异常闭环。",
        f"质量意识完整：能从{spec['focus']}的正确性、性能和可维护性三个维度说明取舍。",
    )
    collaboration_variants = (
        "协作边界明确：会主动同步依赖、风险和需要业务决策的事项，减少信息差造成的返工。",
        "交接意识清楚：会补齐接口说明、运行手册和异常样例，让后续维护不依赖口头传递。",
        "复盘态度诚实：能区分个人可控问题与外部约束，并把可改进项落实到下一轮计划。",
        "跨团队配合顺畅：会提前确认验收口径和上线窗口，减少临近交付时的隐性阻塞。",
        "反馈响应及时：遇到指标偏差会先同步影响范围，再提出带验证方式的调整建议。",
    )
    strengths = [strength_variants[index % len(strength_variants)], evidence_variants[(index + 1) % len(evidence_variants)], collaboration_variants[(index + 2) % len(collaboration_variants)]]
    if level == "高级":
        strengths.append(("带教和机制建设：能把个人经验沉淀为评审清单、监控面板或可复用组件。", "团队影响力较强：能通过技术评审、风险清单和复盘机制帮助成员稳定交付。", "架构判断有依据：面对资源和时间约束，会说明取舍并为关键路径保留回滚方案。")[index % 3])
    weakness_first = (
        "在需求快速变化的早期阶段，容易投入较多时间确认边界；目前用时间盒和最小验收样例控制节奏。",
        "面对信息不完整的需求，前期会倾向于多做几轮核验；目前通过先交付小样例来平衡严谨和速度。",
        "遇到跨团队依赖较多的任务时，容易提前考虑较多异常分支；目前用风险分级和阶段验收控制投入。",
        "在熟悉的技术方向上会追求较完整的工程闭环；目前先按业务优先级交付，再安排质量补强。",
        "对关键指标的定义比较谨慎，初期沟通时间可能偏长；目前用指标草案和评审时间盒提升效率。",
    )
    weakness_second = (
        f"对非核心领域的{('业务流程' if index % 2 == 0 else '硬件细节')}需要继续补课，已通过跟岗、文档和小实验补齐认知。",
        f"对{('现场运维' if index % 2 == 0 else '业务财务口径')}的理解还需加深，已安排真实案例复盘并请教领域同事。",
        f"在{('复杂沟通场景' if index % 2 == 0 else '底层性能调优')}方面仍有提升空间，正在用专项练习和复盘记录补齐。",
        f"对{('陌生行业规则' if index % 2 == 0 else '硬件供应链约束')}的经验有限，已通过旁听评审和小规模验证建立判断框架。",
        f"暂时不把{('表达型汇报' if index % 2 == 0 else '非主线工具')}作为优势，正在用短讲、文档和配对实践持续改善。",
    )
    weaknesses = [weakness_first[index % len(weakness_first)], weakness_second[(index + 2) % len(weakness_second)]]
    return strengths, weaknesses


def _record(index: int, seed: int, level: str, difficulty: str) -> dict[str, Any]:
    spec = ROLE_SPECS[(index // 25) * 4 + (index % 25) % 4]
    education, work, experience = _timeline(index, level, spec)
    projects = _projects(index, level, spec, work)
    groups = spec["groups"]
    flat_skills = _flatten_skills(groups)
    strengths, weaknesses = _strengths_weaknesses(index, level, spec)
    profile = _profile(index, level, spec, experience)
    profile["name"] = NAMES[index]
    summary = _summary(index, level, spec, experience)
    return {
        "candidate_profile": profile,
        "summary": summary,
        "job_intent": spec["role"],
        "role": spec["role"],
        "role_family": spec["family"],
        "domain": spec["domain"],
        "level": level,
        "difficulty": difficulty,
        "skills": flat_skills,
        "skills_grouped": groups,
        "skills_groups": groups,
        "skills_by_category": groups,
        "skill_evidence": _skill_evidence(index, spec, projects, work),
        "education": education[0]["degree"],
        "education_detail": education,
        "work_history": work,
        "projects": projects,
        "experience_years": experience,
        "certifications": ["云原生应用基础证书"] if level == "高级" and index % 2 == 0 else [],
        "preferred_locations": profile["preferred_locations"],
        "expected_salary_k": [8 + index % 4, 13 + index % 8] if level == "初级" else [18 + index % 6, 28 + index % 10] if level == "中级" else [30 + index % 8, 45 + index % 12],
        "strengths": strengths,
        "weaknesses": weaknesses,
        "quality_scenario": {"scenario": SCENES[(index * 2) % len(SCENES)], "constraints": ["既有系统需兼容", "结果需可观测", "上线过程需可回滚"], "evidence_focus": ["职责边界", "技术动作", "量化结果"], "synthetic": True},
        "quality_labels": ["fictional", "timeline_coherent", "evidence_rich", "quantified_projects", "skills_grouped", "non_contact_data"],
        "source": {"synthetic": True, "generator_seed": seed, "no_personal_contact_data": True},
    }


def _resume_text(record: dict[str, Any], style_name: str) -> str:
    profile = record["candidate_profile"]
    lines = [
        f"{profile['name']}｜{record['job_intent']}",
        f"{record['domain']} · {record['level']} · {record['role_family']} · 相关经历 {record['experience_years']:.1f} 年",
        "",
        "职业摘要",
        record["summary"],
        "",
        "求职意向",
        f"目标岗位：{record['job_intent']}；期望地点：{'、'.join(record['preferred_locations'])}；到岗：{profile['availability']}。",
        "",
        "教育背景",
    ]
    for item in record["education_detail"]:
        lines.append(f"{item['start']}—{item['end']}｜{item['institution']}｜{item['degree']}｜{item['major']}")
        lines.append(f"课程与研究：{'、'.join(item['coursework'])}；毕业实践：{item['capstone']}。")
    lines.extend(["", "核心技能"])
    for group, skills in record["skills_grouped"].items():
        lines.append(f"{group}：{'、'.join(skills)}")
    lines.extend(["", "技能证据"])
    for item in record["skill_evidence"][:10]:
        lines.append(f"{item['skill']}（{item['group']}）：{item['evidence']}")
    lines.extend(["", "工作经历"])
    for item in record["work_history"]:
        lines.append(f"{item['start']}—{item['end']}｜{item['company']}｜{item['title']}")
        lines.append(f"职责边界：{item['scope']}。")
        lines.append(f"工作动作：{'；'.join(item['responsibilities'])}。结果：{'；'.join(item['achievements'])}。")
    lines.extend(["", "项目经历"])
    for number, project in enumerate(record["projects"], 1):
        lines.extend([
            f"项目{number}｜{project['name']}｜{project['period']}｜{project['role']}",
            f"背景：{project['background']}",
            f"职责边界：{project['responsibility_boundary']}",
            f"技术动作：{'；'.join(project['technical_actions'])}。",
            f"量化结果：{'；'.join(project['quantified_results'])}。",
        ])
    lines.extend(["", "优势"])
    lines.extend(f"- {item}" for item in record["strengths"])
    lines.extend(["", "合理弱点与改进"])
    lines.extend(f"- {item}" for item in record["weaknesses"])
    lines.extend(["", f"样式：{style_name}", "说明：本简历为虚构合成样本，仅用于简历解析与匹配评测，不对应真实个人。"])
    return "\n".join(lines) + "\n"


def _pdf_font() -> str:
    name = "ResumeCJK"
    if name in pdfmetrics.getRegisteredFontNames():
        return name
    root = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts"
    candidates = [
        Path(os.environ.get("TALENTGRAPH_CJK_FONT_PATH", "")),
        root / "Noto Sans SC (TrueType).otf",
        root / "Deng.ttf",
        root / "simhei.ttf",
        Path("/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            try:
                pdfmetrics.registerFont(TTFont(name, str(candidate)))
                return name
            except Exception:
                continue
    return "Helvetica"


def _write_pdf(path: Path, text: str, layout: dict[str, Any]) -> None:
    font = _pdf_font()
    styles = getSampleStyleSheet()
    body = ParagraphStyle("ResumeBody", parent=styles["BodyText"], fontName=font, fontSize=9.2, leading=14, textColor=colors.HexColor("#25313C"), spaceAfter=3)
    heading = ParagraphStyle("ResumeHeading", parent=body, fontSize=12.5, leading=18, textColor=layout["pdf"], spaceBefore=8, spaceAfter=4)
    title = ParagraphStyle("ResumeTitle", parent=body, fontSize=20, leading=25, textColor=layout["pdf"], spaceAfter=2)
    meta = ParagraphStyle("ResumeMeta", parent=body, fontSize=9.5, leading=14, textColor=colors.HexColor("#52606D"), spaceAfter=8)
    story: list[Any] = []
    for number, raw in enumerate(text.splitlines()):
        line = escape(raw)
        if number == 0:
            story.append(Paragraph(line, title))
        elif number == 1:
            story.append(Paragraph(line, meta))
        elif raw in {"职业摘要", "求职意向", "教育背景", "核心技能", "技能证据", "工作经历", "项目经历", "优势", "合理弱点与改进"}:
            story.extend([HRFlowable(width="100%", thickness=0.6, color=layout["pdf"], spaceBefore=4, spaceAfter=2), Paragraph(line, heading)])
        elif raw.startswith("项目") and "｜" in raw:
            story.append(Paragraph(f"<b>{line}</b>", body))
        elif raw.startswith("- "):
            story.append(Paragraph(f"&#8226; {escape(raw[2:])}", body))
        elif raw:
            story.append(Paragraph(line, body))
        else:
            story.append(Spacer(1, 2))
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=17 * mm, leftMargin=17 * mm, topMargin=15 * mm, bottomMargin=15 * mm, title="虚构中文技术简历")
    doc.build(story, canvasmaker=_stable_canvas)


def _stable_canvas(filename: str, *args: Any, **kwargs: Any) -> pdf_canvas.Canvas:
    """Keep PDF metadata and document IDs reproducible across runs."""
    kwargs["invariant"] = 1
    return pdf_canvas.Canvas(filename, *args, **kwargs)


def _write_docx(path: Path, text: str, layout: dict[str, Any]) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    for number, raw in enumerate(text.splitlines()):
        if not raw:
            document.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.1
        run = paragraph.add_run(raw)
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        if number == 0:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run.bold = True
            run.font.size = Pt(20)
            run.font.color.rgb = RGBColor.from_string(layout["hex"])
        elif raw in {"职业摘要", "求职意向", "教育背景", "核心技能", "技能证据", "工作经历", "项目经历", "优势", "合理弱点与改进"}:
            run.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor.from_string(layout["hex"])
        elif raw.startswith("项目") and "｜" in raw:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(layout["hex"])
        elif raw.startswith("- "):
            paragraph.paragraph_format.left_indent = Inches(0.18)
    document.core_properties.title = "虚构中文技术简历"
    document.core_properties.subject = "Resume Gold 100 benchmark"
    document.core_properties.created = datetime(2020, 1, 1)
    document.core_properties.modified = datetime(2020, 1, 1)
    document.save(path)
    _stabilize_docx(path)


def _stabilize_docx(path: Path) -> None:
    """Normalize ZIP entry timestamps so artifact hashes are repeatable."""
    temp = path.with_name(path.name + ".stable")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for item in source.infolist():
            fixed = zipfile.ZipInfo(item.filename, date_time=(2020, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = item.external_attr
            fixed.create_system = item.create_system
            target.writestr(fixed, source.read(item.filename))
    temp.replace(path)


def _quality_gate(manifest: dict[str, Any], output: Path) -> dict[str, Any]:
    errors: list[str] = []
    records = manifest["records"]
    if len(records) != 100:
        errors.append(f"record_count={len(records)}")
    domain_counts = {domain: sum(item["domain"] == domain for item in records) for domain in DOMAINS}
    if domain_counts != {domain: 25 for domain in DOMAINS}:
        errors.append(f"domain_counts={domain_counts}")
    if len({item["role"] for item in records}) < 12:
        errors.append("role_count<12")
    texts: list[str] = []
    for item in records:
        gold = json.loads((output / item["gold_file"]).read_text(encoding="utf-8"))
        text = (output / item["file_paths"]["txt"]).read_text(encoding="utf-8")
        texts.append(re.sub(r"\s+", "", text))
        minimum = 700 if item["level"] == "初级" else 900
        if len(text.replace("\n", "")) < minimum:
            errors.append(f"{item['id']}:text_short")
        required = ("职业摘要", "求职意向", "教育背景", "核心技能", "工作经历", "项目经历", "优势", "合理弱点与改进")
        if any(section not in text for section in required):
            errors.append(f"{item['id']}:missing_section")
        if not 2 <= len(gold["projects"]) <= 3:
            errors.append(f"{item['id']}:project_count")
        for project in gold["projects"]:
            if not all(project.get(key) for key in ("background", "responsibility_boundary", "technical_actions", "quantified_results")):
                errors.append(f"{item['id']}:project_evidence")
            if not any(re.search(r"\d", result) for result in project["quantified_results"]):
                errors.append(f"{item['id']}:project_metric")
        if len(gold["skill_evidence"]) < 8:
            errors.append(f"{item['id']}:skill_evidence")
        starts = [entry["start"] for entry in gold["work_history"]]
        if starts != sorted(starts) or any(gold["work_history"][pos]["end"] > gold["work_history"][pos + 1]["start"] for pos in range(len(gold["work_history"]) - 1)):
            errors.append(f"{item['id']}:timeline")
        education = gold["education_detail"]
        if education != sorted(education, key=lambda entry: entry["start"]) or any(education[pos]["end"] > education[pos + 1]["start"] for pos in range(len(education) - 1)):
            errors.append(f"{item['id']}:education_timeline")
        if education and starts and max(entry["end"] for entry in education) > starts[0]:
            errors.append(f"{item['id']}:education_work_overlap")
        for kind, artifact in item["artifacts"].items():
            path = output / artifact["path"]
            if not path.exists() or not path.read_bytes():
                errors.append(f"{item['id']}:{kind}_missing")
            if artifact["sha256"] != _sha256(path):
                errors.append(f"{item['id']}:{kind}_hash")
    if len(set(texts)) != 100:
        errors.append("resume_text_not_unique")
    return {"passed": not errors, "errors": errors, "checks": {"records": len(records), "domains": domain_counts, "roles": len({item['role'] for item in records}), "unique_texts": len(set(texts))}}


def generate(output: Path = DEFAULT_OUT, seed: int = SEED, count: int = 100) -> dict[str, Any]:
    """Generate all artifacts and return the manifest."""
    output = Path(output)
    if output.exists():
        shutil.rmtree(output)
    files_dir, gold_dir = output / "files", output / "gold"
    files_dir.mkdir(parents=True)
    gold_dir.mkdir(parents=True)
    levels = _level_sequence(seed)
    difficulties = _difficulty_sequence(seed + 11)
    records: list[dict[str, Any]] = []
    for index in range(count):
        domain_slot = index // 25
        spec = ROLE_SPECS[domain_slot * 4 + (index % 25) % 4]
        level = levels[index]
        difficulty = difficulties[index]
        style = LAYOUTS[index % len(LAYOUTS)]
        gold = _record(index, seed, level, difficulty)
        text = _resume_text(gold, style["name"])
        rid = f"resume_{index + 1:03d}"
        txt_path = files_dir / f"{rid}.txt"
        pdf_path = files_dir / f"{rid}.pdf"
        docx_path = files_dir / f"{rid}.docx"
        txt_path.write_text(text, encoding="utf-8")
        _write_pdf(pdf_path, text, style)
        _write_docx(docx_path, text, style)
        relative = {"txt": f"files/{txt_path.name}", "pdf": f"files/{pdf_path.name}", "docx": f"files/{docx_path.name}"}
        hashes = {kind: _sha256(output / path) for kind, path in relative.items()}
        gold["artifact_paths"] = relative
        gold["artifact_hashes"] = hashes
        gold_path = gold_dir / f"{rid}.json"
        gold_path.write_text(json.dumps(gold, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        artifacts = {kind: {"path": path, "sha256": hashes[kind], "bytes": (output / path).stat().st_size} for kind, path in relative.items()}
        artifacts["gold"] = {"path": f"gold/{gold_path.name}", "sha256": _sha256(gold_path), "bytes": gold_path.stat().st_size}
        split = "calibration" if index < 20 else "regression" if index < 50 else "blind"
        records.append({"id": rid, "filename": pdf_path.name, "format": "可复制中文PDF", "available_formats": ["DOCX", "可复制中文PDF", "TXT"], "file_paths": {**relative, "gold": artifacts["gold"]["path"]}, "file_hashes": {**hashes, "gold": artifacts["gold"]["sha256"]}, "artifacts": artifacts, "domain": spec["domain"], "role": spec["role"], "level": level, "difficulty": difficulty, "layout_style": style["name"], "split": split, "gold_file": artifacts["gold"]["path"], "sha256": hashes["pdf"], "bytes": artifacts["pdf"]["bytes"], "quality_labels": gold["quality_labels"]})
    manifest = {"dataset": "resume_gold_100", "schema_version": "2.0", "count": count, "synthetic": True, "description": "100份虚构匿名中文技术简历；每人同时提供DOCX、可复制中文PDF、TXT和增强gold JSON。", "formats": ["DOCX", "可复制中文PDF", "TXT", "增强gold JSON"], "layout_styles": [item["name"] for item in LAYOUTS], "records": records}
    manifest["quality_gate"] = _quality_gate(manifest, output) if count == 100 else {"passed": True, "errors": [], "checks": {"records": count}}
    if not manifest["quality_gate"]["passed"]:
        raise RuntimeError("resume gold quality gate failed: " + "; ".join(manifest["quality_gate"]["errors"][:8]))
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--seed", type=int, default=SEED)
    parser.add_argument("--count", type=int, default=100)
    args = parser.parse_args()
    result = generate(args.output, args.seed, args.count)
    print(json.dumps({"dataset": result["dataset"], "count": result["count"], "quality_gate": result["quality_gate"]}, ensure_ascii=False, indent=2))
