from __future__ import annotations

from html import escape
from io import BytesIO
import os
from pathlib import Path
from typing import Any


BLUE = "2F6DF4"
BLUE_DARK = "17365D"
BLUE_LIGHT = "EAF2FF"
BLUE_LINE = "BDD1F5"
INK = "20324A"
MUTED = "60758F"
WORD_CJK_FONT = "Microsoft YaHei"
PDF_CJK_FONT = "TalentGraphCJK"


def export_match_report(report: Any, profile: Any, output_format: str) -> tuple[bytes, str, str]:
    if output_format == "docx":
        return _build_docx(report, profile), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
    return _build_pdf(report, profile), "application/pdf", "pdf"


def _profile_value(profile: Any, key: str, default: Any = None) -> Any:
    if profile is None:
        return default
    if hasattr(profile, key):
        return getattr(profile, key)
    if isinstance(profile, dict):
        return profile.get(key, default)
    return default


def _report_sections(report: Any, profile: Any) -> list[tuple[str, list[str]]]:
    strengths = [
        f"{item.get('skill', '能力优势')}：{item.get('summary') or item.get('evidence') or '已有候选人证据支持'}"
        for item in (report.strengths or [])
    ]
    gaps = [
        f"{item.get('skill', '能力缺口')}：{item.get('suggestion') or item.get('recommended_action') or '建议在终面进一步核验'}"
        for item in (report.gaps or [])
    ]
    learning = [
        f"{item.get('step') or item.get('title') or f'阶段 {index + 1}'}：{item.get('action') or item.get('description') or str(item)}"
        for index, item in enumerate(report.learning_path or [])
    ]
    dimensions = [
        f"{item.get('name', '能力维度')}：{item.get('score', 0)} 分"
        for item in (_profile_value(profile, "competency_dimensions", []) or [])
    ]
    return [
        ("招聘结论", [
            report.executive_summary,
            f"综合匹配度：{report.score}/100",
            f"推荐结论：{report.interview_recommendation or _level_label(report.level)}",
            f"证据覆盖率：{round(report.evidence_coverage * 100)}%",
        ]),
        ("能力人格画像", [
            f"{_profile_value(profile, 'archetype_title', '待完成能力画像')}（{_profile_value(profile, 'archetype_code', 'N/A')}）",
            _profile_value(profile, "archetype_narrative", "") or _profile_value(profile, "summary", ""),
            *dimensions,
        ]),
        ("优势证据", strengths or ["暂无明确优势证据。"]),
        ("能力差距", gaps or ["未识别到关键能力缺口。"]),
        ("风险提示", list(report.risks or []) or ["当前未识别到显著风险。"]),
        ("终面追问", list(report.follow_up_questions or []) or ["建议围绕岗位真实场景继续核验。"]),
        ("补强路径", learning or ["暂无补强路径。"]),
    ]


def _level_label(level: str) -> str:
    return {"high": "建议优先推进", "medium": "建议进入终面核验", "low": "建议暂缓推进"}.get(level, level)


def _candidate_label(report: Any, profile: Any) -> str:
    return str(
        _profile_value(profile, "name", "")
        or _profile_value(profile, "candidate_name", "")
        or report.candidate_id
    )


def _job_title(report: Any) -> str:
    return str(report.job_snapshot.get("job_title") or report.target_job_id)


def _build_docx(report: Any, profile: Any) -> bytes:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt, RGBColor

    def set_style_font(style, size: float, color: str, *, bold: bool = False) -> None:
        style.font.name = WORD_CJK_FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{key}"), WORD_CJK_FONT)

    def set_run_font(run, size: float, color: str, *, bold: bool = False) -> None:
        run.font.name = WORD_CJK_FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{key}"), WORD_CJK_FONT)

    def shade_paragraph(paragraph, fill: str, accent: str) -> None:
        properties = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "24")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), accent)
        borders.append(left)
        properties.append(borders)

    def add_page_number(paragraph) -> None:
        set_run_font(paragraph.add_run("第 "), 8.5, MUTED)
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        value = OxmlElement("w:t")
        value.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend([begin, instruction, separate, value, end])
        set_run_font(run, 8.5, MUTED)
        set_run_font(paragraph.add_run(" 页"), 8.5, MUTED)

    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(19)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    section.header_distance = Mm(9)
    section.footer_distance = Mm(9)

    normal = document.styles["Normal"]
    set_style_font(normal, 10.5, INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading = document.styles["Heading 1"]
    set_style_font(heading, 14, BLUE_DARK, bold=True)
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.keep_with_next = True

    bullets = document.styles["List Bullet"]
    set_style_font(bullets, 10.5, INK)
    bullets.paragraph_format.left_indent = Inches(0.38)
    bullets.paragraph_format.first_line_indent = Inches(-0.18)
    bullets.paragraph_format.space_after = Pt(5)
    bullets.paragraph_format.line_spacing = 1.2

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_run_font(header.add_run("TALENTGRAPH EVO  |  MATCH INTELLIGENCE"), 8.5, BLUE, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    set_run_font(footer.add_run("TalentGraph Evo · 数字人才匹配报告  |  "), 8.5, MUTED)
    add_page_number(footer)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(5)
    set_run_font(kicker.add_run("DIGITAL TALENT MATCH REPORT"), 9, BLUE, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(7)
    set_run_font(title.add_run("企业招聘匹配报告"), 22, BLUE_DARK, bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_run_font(
        subtitle.add_run(f"候选人：{_candidate_label(report, profile)}    目标岗位：{_job_title(report)}"),
        10.5,
        MUTED,
    )

    recommendation = report.interview_recommendation or _level_label(report.level)
    score = document.add_paragraph()
    score.paragraph_format.left_indent = Mm(4)
    score.paragraph_format.right_indent = Mm(4)
    score.paragraph_format.space_before = Pt(0)
    score.paragraph_format.space_after = Pt(9)
    score.paragraph_format.line_spacing = 1.15
    set_run_font(score.add_run("综合匹配\n"), 9.5, MUTED, bold=True)
    set_run_font(score.add_run(f"{report.score} / 100"), 28, BLUE, bold=True)
    set_run_font(score.add_run(f"\n{recommendation}"), 11, BLUE_DARK, bold=True)
    shade_paragraph(score, BLUE_LIGHT, BLUE)

    for section_heading, items in _report_sections(report, profile):
        document.add_heading(section_heading, level=1)
        for item in items:
            text = str(item).strip()
            if not text:
                continue
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.keep_together = True
            set_run_font(paragraph.add_run(text), 10.5, INK)

    disclaimer = document.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(12)
    disclaimer.paragraph_format.space_after = Pt(0)
    set_run_font(
        disclaimer.add_run("本报告由岗位图谱、候选人档案、深度对话与互动任务证据共同生成，仅作为招聘决策辅助。"),
        8.5,
        MUTED,
    )

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf_font_name() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont

    if PDF_CJK_FONT in pdfmetrics.getRegisteredFontNames():
        return PDF_CJK_FONT

    configured = os.getenv("TALENTGRAPH_CJK_FONT_PATH", "").strip()
    candidates = [
        configured,
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\NotoSansSC-VF.ttf",
        r"C:\Windows\Fonts\simhei.ttf",
        "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/truetype/arphic/uming.ttc",
        "/usr/share/fonts/truetype/arphic/ukai.ttc",
    ]
    for candidate in candidates:
        if not candidate or not Path(candidate).is_file():
            continue
        try:
            font = TTFont(PDF_CJK_FONT, candidate)
            # Some CJK fallback fonts omit Latin glyphs and corrupt mixed text.
            required = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789<>企业招聘匹配报告"
            if any(not font.face.charToGlyph.get(ord(char)) for char in required):
                continue
            pdfmetrics.registerFont(font)
            return PDF_CJK_FONT
        except Exception:
            continue

    fallback = "STSong-Light"
    if fallback not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(fallback))
    return fallback


def _pdf_text(value: Any) -> str:
    return escape(str(value or ""), quote=False).replace("\n", "<br/>")


def _build_pdf(report: Any, profile: Any) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_name = _pdf_font_name()
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=22 * mm,
        bottomMargin=19 * mm,
        title="TalentGraph Evo 企业招聘匹配报告",
        author="TalentGraph Evo",
        subject="数字人才岗位匹配报告",
    )
    styles = getSampleStyleSheet()
    kicker_style = ParagraphStyle(
        "ChineseKicker",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=8.5,
        leading=12,
        spaceAfter=4,
        textColor=colors.HexColor(f"#{BLUE}"),
        wordWrap="CJK",
    )
    title_style = ParagraphStyle(
        "ChineseTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=21,
        leading=29,
        alignment=TA_LEFT,
        spaceAfter=6,
        textColor=colors.HexColor(f"#{BLUE_DARK}"),
        wordWrap="CJK",
    )
    heading_style = ParagraphStyle(
        "ChineseHeading",
        parent=styles["Heading2"],
        fontName=font_name,
        fontSize=13,
        leading=20,
        spaceBefore=11,
        spaceAfter=5,
        keepWithNext=True,
        textColor=colors.HexColor(f"#{BLUE_DARK}"),
        wordWrap="CJK",
    )
    body_style = ParagraphStyle(
        "ChineseBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=10.5,
        leading=17,
        leftIndent=5 * mm,
        firstLineIndent=-4 * mm,
        spaceAfter=3.5,
        textColor=colors.HexColor(f"#{INK}"),
        wordWrap="CJK",
    )
    small_style = ParagraphStyle(
        "ChineseSmall",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=8.5,
        leading=13,
        textColor=colors.HexColor(f"#{MUTED}"),
        wordWrap="CJK",
    )
    score_style = ParagraphStyle(
        "ChineseScore",
        parent=title_style,
        fontName=font_name,
        fontSize=27,
        leading=33,
        alignment=TA_CENTER,
        textColor=colors.HexColor(f"#{BLUE}"),
    )
    recommendation_style = ParagraphStyle(
        "ChineseRecommendation",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=11,
        leading=18,
        textColor=colors.HexColor(f"#{BLUE_DARK}"),
        wordWrap="CJK",
    )

    def draw_page(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont(font_name, 8)
        canvas.setFillColor(colors.HexColor(f"#{BLUE}"))
        canvas.drawString(20 * mm, A4[1] - 11 * mm, "TALENTGRAPH EVO  |  MATCH INTELLIGENCE")
        canvas.setStrokeColor(colors.HexColor(f"#{BLUE_LINE}"))
        canvas.setLineWidth(0.6)
        canvas.line(20 * mm, A4[1] - 13 * mm, A4[0] - 20 * mm, A4[1] - 13 * mm)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.drawString(20 * mm, 10 * mm, "TalentGraph Evo · 数字人才匹配报告")
        canvas.drawRightString(A4[0] - 20 * mm, 10 * mm, f"第 {doc.page} 页")
        canvas.restoreState()

    recommendation = report.interview_recommendation or _level_label(report.level)
    story = [
        Paragraph("DIGITAL TALENT MATCH REPORT", kicker_style),
        Paragraph("企业招聘匹配报告", title_style),
        Paragraph(
            _pdf_text(f"候选人：{_candidate_label(report, profile)}    目标岗位：{_job_title(report)}"),
            small_style,
        ),
        Spacer(1, 6 * mm),
        Table(
            [[
                Paragraph(_pdf_text(f"{report.score} / 100"), score_style),
                Paragraph(_pdf_text(recommendation), recommendation_style),
            ]],
            colWidths=[50 * mm, 120 * mm],
            style=TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{BLUE_LIGHT}")),
                ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor(f"#{BLUE_LINE}")),
                ("LINEBEFORE", (0, 0), (0, 0), 3, colors.HexColor(f"#{BLUE}")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 12),
                ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                ("TOPPADDING", (0, 0), (-1, -1), 11),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 11),
            ]),
        ),
    ]
    for section_heading, items in _report_sections(report, profile):
        story.append(Paragraph(_pdf_text(section_heading), heading_style))
        for item in items:
            text = str(item).strip()
            if text:
                story.append(Paragraph(f"•　{_pdf_text(text)}", body_style))
    story.extend([
        Spacer(1, 5 * mm),
        Paragraph(
            "本报告由岗位图谱、候选人档案、深度对话与互动任务证据共同生成，仅作为招聘决策辅助。",
            small_style,
        ),
    ])
    document.build(story, onFirstPage=draw_page, onLaterPages=draw_page)
    return buffer.getvalue()
